import streamlit as st
import pandas as pd
from google import genai
import matplotlib.pyplot as plt
import seaborn as sns
import hashlib
import hmac
import io
from io import BytesIO
import json
import secrets
import sqlite3
import re
import uuid
import itertools
import math
from contextlib import closing
from datetime import datetime
from pathlib import Path

import numpy as np
from PIL import Image, UnidentifiedImageError
from google.genai import types
from scipy import stats
from ui_theme import apply_theme

st.set_page_config(
    page_title="Market Research AI Data Analyst Workspace",
    page_icon=str(Path(__file__).resolve().with_name("LogoStatlabChat_transparant.png")),
    layout="wide",
    initial_sidebar_state="auto",
)
apply_theme()


# --- Modern dark UI override: charcoal surfaces + soft cyan accent ---
# This layer changes presentation only; the existing application logic remains intact.
st.markdown(
    """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&display=swap');

    :root {
        --sl-bg: #0b0f14;
        --sl-sidebar: #0e141c;
        --sl-panel: #121a24;
        --sl-panel-hover: #172330;
        --sl-border: rgba(157, 190, 219, .14);
        --sl-text: #e8eef6;
        --sl-muted: #91a0b2;
        --sl-blue: #82d0ff;
        --sl-blue-strong: #4eafea;
    }

    html, body, [class*="css"] {
        font-family: 'DM Sans', sans-serif;
    }
    .stApp {
        background:
            radial-gradient(ellipse at 55% -15%, rgba(61, 144, 199, .09), transparent 42%),
            var(--sl-bg);
        color: var(--sl-text);
    }
    [data-testid="stHeader"] {
        background: rgba(11, 15, 20, .88);
    }
    [data-testid="stSidebar"] {
        background: var(--sl-sidebar);
        border-right: 1px solid var(--sl-border);
    }
    [data-testid="stSidebar"] > div {
        padding-top: 1.1rem;
    }
    [data-testid="stSidebar"] [data-testid="stImage"] {
        padding: 4px 0 8px;
    }
    [data-testid="stSidebar"] [data-testid="stCaptionContainer"] {
        color: #8fa2b5;
        font-size: 12px;
    }
    .main .block-container {
        max-width: 1180px;
        padding-top: 2rem;
        padding-bottom: 4rem;
    }
    h1, h2, h3, h4 {
        color: #f0f6fc;
        letter-spacing: -.035em;
    }
    h1 { font-weight: 700; }
    p, li, label, [data-testid="stMarkdownContainer"] {
        color: var(--sl-text);
    }
    [data-testid="stCaptionContainer"], .stCaption {
        color: var(--sl-muted) !important;
    }

    /* Sidebar navigation buttons */
    [data-testid="stSidebar"] .stButton > button {
        text-align: left;
        min-height: 42px;
        border-radius: 11px;
        border: 1px solid transparent;
        background: transparent;
        color: #b8c7d6;
        font-weight: 500;
        transition: background .18s ease, border-color .18s ease, color .18s ease;
    }
    [data-testid="stSidebar"] .stButton > button:hover {
        background: #15212d;
        border-color: var(--sl-border);
        color: #f4f8fc;
    }
    [data-testid="stSidebar"] .stButton > button[kind="primary"] {
        background: linear-gradient(135deg, rgba(130,208,255,.18), rgba(78,175,234,.08));
        border-color: rgba(130,208,255,.27);
        color: #bce7ff;
        box-shadow: inset 3px 0 0 var(--sl-blue);
    }
    [data-testid="stSidebar"] hr {
        border-color: var(--sl-border);
    }
    .sidebar-app-brand {
        margin: 18px 8px 4px;
        color: #7f93a7;
        font-size: 12px;
        line-height: 1.6;
        letter-spacing: .03em;
    }

    /* Inputs and forms */
    input, textarea, [data-baseweb="select"] > div {
        border-radius: 11px !important;
    }
    [data-testid="stTextInput"] input,
    [data-testid="stTextArea"] textarea,
    [data-testid="stNumberInput"] input,
    [data-testid="stDateInput"] input,
    [data-testid="stFileUploader"] section {
        background: #101720 !important;
        color: var(--sl-text) !important;
        border-color: var(--sl-border) !important;
    }
    [data-testid="stTextInput"] input:focus,
    [data-testid="stTextArea"] textarea:focus {
        border-color: rgba(130,208,255,.65) !important;
        box-shadow: 0 0 0 1px rgba(130,208,255,.18) !important;
    }
    .stButton > button, [data-testid="stFormSubmitButton"] button,
    [data-testid="stDownloadButton"] button {
        min-height: 40px;
        border-radius: 11px;
        border: 1px solid var(--sl-border);
        background: #141d27;
        color: #dbe7f2;
        font-weight: 600;
        transition: all .18s ease;
    }
    .stButton > button:hover, [data-testid="stDownloadButton"] button:hover {
        border-color: rgba(130,208,255,.42);
        background: #1a2937;
        color: #fff;
    }
    button[kind="primary"], [data-testid="stFormSubmitButton"] button[kind="primary"] {
        background: linear-gradient(135deg, #a1e0ff, #65bdf1) !important;
        border: 0 !important;
        color: #07131d !important;
        font-weight: 700 !important;
    }

    /* ChatGPT-like chat messages and composer */
    [data-testid="stChatMessage"] {
        background: rgba(255,255,255,.018);
        border: 1px solid rgba(157,190,219,.08);
        border-radius: 16px;
        padding: 15px 17px;
        margin-bottom: 12px;
    }
    [data-testid="stChatMessage"]:has([data-testid="chatAvatarIcon-assistant"]) {
        background: rgba(130,208,255,.025);
    }
    [data-testid="stChatInput"] {
        border: 1px solid rgba(130,208,255,.27) !important;
        border-radius: 18px !important;
        background: #111923 !important;
        box-shadow: 0 12px 38px rgba(0,0,0,.20);
    }
    [data-testid="stChatInput"]:focus-within {
        border-color: rgba(130,208,255,.62) !important;
        box-shadow: 0 0 0 1px rgba(130,208,255,.12), 0 12px 38px rgba(0,0,0,.2);
    }
    [data-testid="stChatInput"] textarea {
        color: #eaf3fb !important;
    }
    [data-testid="stChatInput"] textarea::placeholder {
        color: #788b9e !important;
    }

    /* Tables, expanders, cards, upload panels */
    [data-testid="stDataFrame"], [data-testid="stTable"] {
        border: 1px solid var(--sl-border);
        border-radius: 12px;
        overflow: hidden;
    }
    [data-testid="stExpander"] {
        border: 1px solid var(--sl-border);
        border-radius: 13px;
        background: rgba(255,255,255,.015);
    }
    [data-testid="stAlert"] {
        border-radius: 12px;
    }
    [data-testid="stFileUploader"] {
        border-radius: 14px;
    }
    .login-brand {
        text-align: center;
        padding: 28px 0 18px;
    }
    .login-brand h1 {
        font-size: clamp(25px, 3vw, 36px);
        letter-spacing: -.045em;
        margin-bottom: 8px;
        background: linear-gradient(100deg, #f0f7ff 20%, #8bd5ff 80%);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
    }
    .login-brand p {
        color: var(--sl-muted);
        font-size: 14px;
    }
    [data-testid="stForm"] {
        background: rgba(255,255,255,.018);
        border: 1px solid var(--sl-border);
        border-radius: 16px;
        padding: 22px;
    }
    footer { visibility: hidden; }
    #MainMenu { visibility: hidden; }

    @media (max-width: 700px) {
        .main .block-container {
            padding-top: 1rem;
            padding-left: 1rem;
            padding-right: 1rem;
        }
    }
    </style>
    """,
    unsafe_allow_html=True,
)

USER_DB_PATH = Path(__file__).resolve().parent / "user" / "user.db"
DATA_PROJECT_DIR = Path(__file__).resolve().parent / "data_project"
CHAT_HISTORY_DIR = Path(__file__).resolve().parent / "history_chat"
PHOTO_HISTORY_DIR = Path(__file__).resolve().parent / "history_foto"
STATISTICS_HISTORY_DIR = Path(__file__).resolve().parent / "history_statistik"
LOGO_PATH = Path(__file__).resolve().with_name("LogoStatlabChat_transparant.png")
PASSWORD_HASH_ITERATIONS = 600_000
USER_ROLES = ("admin", "statistician", "researcher")
DEMO_USERS = (
    ("demoadmin@lab.chat", "12345", "admin"),
    ("demodata@lab.chat", "12345", "statistician"),
    ("demore@lab.chat", "12345", "researcher"),
)


def hash_password(password: str) -> str:
    salt = secrets.token_bytes(16)
    password_hash = hashlib.pbkdf2_hmac(
        "sha256",
        password.encode("utf-8"),
        salt,
        PASSWORD_HASH_ITERATIONS,
    ).hex()
    return f"pbkdf2_sha256${PASSWORD_HASH_ITERATIONS}${salt.hex()}${password_hash}"


def verify_password(password: str, stored_hash: str) -> bool:
    try:
        algorithm, iterations, salt_hex, expected_hash = stored_hash.split("$")
        if algorithm != "pbkdf2_sha256":
            return False
        actual_hash = hashlib.pbkdf2_hmac(
            "sha256",
            password.encode("utf-8"),
            bytes.fromhex(salt_hex),
            int(iterations),
        ).hex()
    except (ValueError, TypeError):
        return False
    return hmac.compare_digest(actual_hash, expected_hash)


def initialize_user_database() -> None:
    USER_DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    with closing(sqlite3.connect(USER_DB_PATH)) as connection:
        with connection:
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS users (
                    username TEXT PRIMARY KEY,
                    password_hash TEXT NOT NULL,
                    role TEXT NOT NULL DEFAULT 'researcher'
                        CHECK (role IN ('admin', 'statistician', 'researcher'))
                )
                """
            )
            columns = {
                row[1] for row in connection.execute("PRAGMA table_info(users)")
            }
            if "role" not in columns:
                connection.execute(
                    """
                    ALTER TABLE users
                    ADD COLUMN role TEXT NOT NULL DEFAULT 'researcher'
                        CHECK (role IN ('admin', 'statistician', 'researcher'))
                    """
                )

            for username, password, role in DEMO_USERS:
                user_exists = connection.execute(
                    "SELECT 1 FROM users WHERE username = ?",
                    (username,),
                ).fetchone()
                if user_exists is None:
                    connection.execute(
                        """
                        INSERT INTO users (username, password_hash, role)
                        VALUES (?, ?, ?)
                        """,
                        (username, hash_password(password), role),
                    )


def save_user(username: str, password: str, role: str) -> None:
    if role not in USER_ROLES:
        raise ValueError("Tipe user tidak valid.")

    with closing(sqlite3.connect(USER_DB_PATH)) as connection:
        with connection:
            connection.execute(
                "INSERT INTO users (username, password_hash, role) VALUES (?, ?, ?)",
                (username, hash_password(password), role),
            )


def list_users() -> pd.DataFrame:
    with closing(sqlite3.connect(USER_DB_PATH)) as connection:
        return pd.read_sql_query(
            "SELECT username, role FROM users ORDER BY username COLLATE NOCASE",
            connection,
        )


def update_user_role(
    username: str,
    new_role: str,
    actor_username: str,
) -> None:
    if new_role not in USER_ROLES:
        raise ValueError("Role pengguna tidak valid.")
    if username == actor_username and new_role != "admin":
        raise ValueError("Admin yang sedang login tidak dapat menurunkan role sendiri.")

    with closing(sqlite3.connect(USER_DB_PATH)) as connection:
        connection.execute("BEGIN IMMEDIATE")
        actor = connection.execute(
            "SELECT role FROM users WHERE username = ?",
            (actor_username,),
        ).fetchone()
        if actor is None or actor[0] != "admin":
            raise PermissionError("Hanya admin terdaftar yang dapat mengubah role pengguna.")
        user = connection.execute(
            "SELECT role FROM users WHERE username = ?",
            (username,),
        ).fetchone()
        if user is None:
            raise FileNotFoundError("Pengguna tidak ditemukan.")
        if user[0] == "admin" and new_role != "admin":
            admin_count = connection.execute(
                "SELECT COUNT(*) FROM users WHERE role = 'admin'"
            ).fetchone()[0]
            if admin_count <= 1:
                raise ValueError("Role admin terakhir tidak dapat diubah.")
        connection.execute(
            "UPDATE users SET role = ? WHERE username = ?",
            (new_role, username),
        )
        connection.commit()


def delete_user(
    username: str,
    actor_username: str,
) -> None:
    if username == actor_username:
        raise ValueError("Admin yang sedang login tidak dapat menghapus akunnya sendiri.")

    with closing(sqlite3.connect(USER_DB_PATH)) as connection:
        connection.execute("BEGIN IMMEDIATE")
        actor = connection.execute(
            "SELECT role FROM users WHERE username = ?",
            (actor_username,),
        ).fetchone()
        if actor is None or actor[0] != "admin":
            raise PermissionError("Hanya admin terdaftar yang dapat menghapus pengguna.")
        user = connection.execute(
            "SELECT role FROM users WHERE username = ?",
            (username,),
        ).fetchone()
        if user is None:
            raise FileNotFoundError("Pengguna tidak ditemukan.")
        if user[0] == "admin":
            admin_count = connection.execute(
                "SELECT COUNT(*) FROM users WHERE role = 'admin'"
            ).fetchone()[0]
            if admin_count <= 1:
                raise ValueError("Admin terakhir tidak dapat dihapus.")
        connection.execute("DELETE FROM users WHERE username = ?", (username,))
        connection.commit()


def authenticate_user(username: str, password: str) -> str | None:
    with closing(sqlite3.connect(USER_DB_PATH)) as connection:
        row = connection.execute(
            "SELECT password_hash, role FROM users WHERE username = ?",
            (username,),
        ).fetchone()
    if row is not None and verify_password(password, row[0]):
        return row[1]
    return None


def save_project_database(
    dataframe: pd.DataFrame,
    source_filename: str,
    raw_dataframe: pd.DataFrame,
    description_dataframe: pd.DataFrame,
) -> tuple[Path, Path]:
    source_stem = Path(source_filename).stem
    safe_stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", source_stem).strip(" .")
    if not safe_stem:
        raise ValueError("Nama file sumber tidak valid untuk nama database.")

    DATA_PROJECT_DIR.mkdir(parents=True, exist_ok=True)
    database_path = DATA_PROJECT_DIR / f"{safe_stem}.db"
    description_path = DATA_PROJECT_DIR / f"{safe_stem}_desc.db"
    now = datetime.now().astimezone()
    submitted_at = (
        get_project_database_timestamps(database_path)[0]
        if database_path.exists()
        else now
    )
    with closing(sqlite3.connect(database_path)) as connection:
        with connection:
            raw_dataframe.to_sql("raw_data", connection, if_exists="replace", index=False)
            dataframe.to_sql("survey_data", connection, if_exists="replace", index=False)
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS project_metadata (
                    id INTEGER PRIMARY KEY CHECK (id = 1),
                    submitted_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                )
                """
            )
            connection.execute(
                """
                INSERT OR REPLACE INTO project_metadata (id, submitted_at, updated_at)
                VALUES (1, ?, ?)
                """,
                (
                    submitted_at.isoformat(timespec="seconds"),
                    now.isoformat(timespec="seconds"),
                ),
            )
    with closing(sqlite3.connect(description_path)) as connection:
        with connection:
            description_dataframe.to_sql(
                "question_descriptions",
                connection,
                if_exists="replace",
                index=False,
            )
    return database_path, description_path


def list_project_databases() -> list[Path]:
    if not DATA_PROJECT_DIR.exists():
        return []
    return sorted(
        (
            database
            for database in DATA_PROJECT_DIR.glob("*.db")
            if not database.stem.endswith("_desc")
        ),
        key=lambda path: path.name.casefold(),
    )


def get_project_database_timestamps(
    database_path: Path,
) -> tuple[datetime, datetime]:
    file_stat = database_path.stat()
    submitted_at = datetime.fromtimestamp(file_stat.st_ctime).astimezone()
    updated_at = datetime.fromtimestamp(file_stat.st_mtime).astimezone()

    with closing(sqlite3.connect(database_path)) as connection:
        metadata_table = connection.execute(
            """
            SELECT 1 FROM sqlite_master
            WHERE type = 'table' AND name = 'project_metadata'
            """
        ).fetchone()
        if metadata_table is not None:
            row = connection.execute(
                "SELECT submitted_at, updated_at FROM project_metadata WHERE id = 1"
            ).fetchone()
            if row is not None:
                submitted_at = datetime.fromisoformat(row[0]).astimezone()
                updated_at = datetime.fromisoformat(row[1]).astimezone()
    return submitted_at, updated_at


def delete_project_database(database_name: str) -> None:
    matching_database = next(
        (
            path
            for path in list_project_databases()
            if path.name == database_name
        ),
        None,
    )
    if matching_database is None:
        raise FileNotFoundError("Database proyek tidak ditemukan.")

    project_directory = DATA_PROJECT_DIR.resolve()
    database_path = matching_database.resolve()
    description_path = matching_database.with_name(
        f"{matching_database.stem}_desc.db"
    ).resolve()
    if (
        database_path.parent != project_directory
        or description_path.parent != project_directory
    ):
        raise ValueError("Lokasi database proyek tidak valid.")

    if description_path.exists():
        description_path.unlink()
    database_path.unlink()


def load_project_database(database_path: Path) -> pd.DataFrame:
    with closing(sqlite3.connect(database_path)) as connection:
        return pd.read_sql_query("SELECT * FROM survey_data", connection)


def load_project_raw_database(database_path: Path) -> pd.DataFrame:
    with closing(sqlite3.connect(database_path)) as connection:
        return pd.read_sql_query("SELECT * FROM raw_data", connection)


def load_project_descriptions(database_path: Path) -> pd.DataFrame:
    description_path = database_path.with_name(f"{database_path.stem}_desc.db")
    with closing(sqlite3.connect(description_path)) as connection:
        columns = {
            row[1]
            for row in connection.execute(
                "PRAGMA table_info(question_descriptions)"
            )
        }
        description_column = (
            "deskripsi_kolom"
            if "deskripsi_kolom" in columns
            else "'' AS deskripsi_kolom"
        )
        return pd.read_sql_query(
            "SELECT no_kues, nama_kolom, "
            f"{description_column}, format_jawaban "
            "FROM question_descriptions ORDER BY no_kues",
            connection,
        )


def get_chat_history_path(username: str) -> Path:
    username_hash = hashlib.sha256(username.encode("utf-8")).hexdigest()
    return CHAT_HISTORY_DIR / f"{username_hash}.json"


def load_chat_histories(username: str) -> dict[str, list[dict[str, object]]]:
    history_path = get_chat_history_path(username)
    if not history_path.exists():
        return {}

    document = json.loads(history_path.read_text(encoding="utf-8"))
    if not isinstance(document, dict) or not isinstance(document.get("conversations"), dict):
        raise ValueError("Format file riwayat chat tidak valid.")

    conversations = document["conversations"]
    for database_name, messages in conversations.items():
        if not isinstance(database_name, str) or not isinstance(messages, list):
            raise ValueError("Format percakapan pada riwayat chat tidak valid.")
        for message in messages:
            if (
                not isinstance(message, dict)
                or message.get("role") not in {"user", "assistant"}
                or not isinstance(message.get("content"), str)
            ):
                raise ValueError("Format pesan pada riwayat chat tidak valid.")
    return conversations


def save_chat_histories(
    username: str,
    conversations: dict[str, list[dict[str, object]]],
) -> None:
    CHAT_HISTORY_DIR.mkdir(parents=True, exist_ok=True)
    history_path = get_chat_history_path(username)
    temporary_path = history_path.with_name(
        f"{history_path.stem}.{secrets.token_hex(8)}.tmp"
    )
    try:
        temporary_path.write_text(
            json.dumps(
                {"username": username, "conversations": conversations},
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        temporary_path.replace(history_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def get_statistics_history_directory(username: str) -> Path:
    username_hash = hashlib.sha256(username.encode("utf-8")).hexdigest()
    return STATISTICS_HISTORY_DIR / username_hash


def get_statistics_history_path(username: str) -> Path:
    return get_statistics_history_directory(username) / "history.json"


def get_statistics_dataset_path(username: str, conversation_id: str) -> Path:
    if not re.fullmatch(r"[a-f0-9]{32}", conversation_id):
        raise ValueError("ID percakapan statistik tidak valid.")
    return get_statistics_history_directory(username) / f"{conversation_id}.csv"


def load_statistics_histories(username: str) -> list[dict[str, object]]:
    history_path = get_statistics_history_path(username)
    if not history_path.exists():
        return []

    document = json.loads(history_path.read_text(encoding="utf-8"))
    if (
        not isinstance(document, dict)
        or not isinstance(document.get("conversations"), list)
    ):
        raise ValueError("Format file riwayat analisis statistik tidak valid.")

    conversations = document["conversations"]
    for conversation in conversations:
        if (
            not isinstance(conversation, dict)
            or not isinstance(conversation.get("id"), str)
            or not re.fullmatch(r"[a-f0-9]{32}", conversation["id"])
            or not isinstance(conversation.get("title"), str)
            or not isinstance(conversation.get("created_at"), str)
            or not isinstance(conversation.get("updated_at"), str)
            or (
                conversation.get("dataset_filename") is not None
                and conversation.get("dataset_filename")
                != f"{conversation['id']}.csv"
            )
            or not isinstance(conversation.get("messages"), list)
        ):
            raise ValueError("Entri percakapan analisis statistik tidak valid.")
        for message in conversation["messages"]:
            if (
                not isinstance(message, dict)
                or message.get("role") not in {"user", "assistant"}
                or not isinstance(message.get("content"), str)
            ):
                raise ValueError("Pesan riwayat analisis statistik tidak valid.")
    return conversations


def save_statistics_histories(
    username: str,
    conversations: list[dict[str, object]],
) -> None:
    history_directory = get_statistics_history_directory(username)
    history_directory.mkdir(parents=True, exist_ok=True)
    history_path = get_statistics_history_path(username)
    temporary_path = history_path.with_name(
        f"{history_path.stem}.{secrets.token_hex(8)}.tmp"
    )
    try:
        temporary_path.write_text(
            json.dumps(
                {"username": username, "conversations": conversations},
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        temporary_path.replace(history_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def delete_statistics_conversation(
    username: str,
    conversations: list[dict[str, object]],
    conversation_id: str,
) -> list[dict[str, object]]:
    dataset_path = get_statistics_dataset_path(username, conversation_id)
    remaining_conversations = [
        conversation
        for conversation in conversations
        if conversation["id"] != conversation_id
    ]
    if len(remaining_conversations) == len(conversations):
        raise FileNotFoundError("Percakapan analisis statistik tidak ditemukan.")

    save_statistics_histories(username, remaining_conversations)
    if dataset_path.exists():
        dataset_path.unlink()
    return remaining_conversations


def analyze_dataframe_automatically(dataframe: pd.DataFrame) -> dict[str, object]:
    numeric_columns = dataframe.select_dtypes(include="number").columns.tolist()
    categorical_columns = [
        column
        for column in dataframe.columns
        if column not in numeric_columns
        and dataframe[column].nunique(dropna=True) <= 100
    ]
    numeric_data = dataframe[numeric_columns].replace(
        [np.inf, -np.inf],
        np.nan,
    )

    column_profile = pd.DataFrame(
        {
            "Kolom": dataframe.columns,
            "Tipe data": [str(dataframe[column].dtype) for column in dataframe.columns],
            "Jumlah data": [int(dataframe[column].count()) for column in dataframe.columns],
            "Nilai kosong": [int(dataframe[column].isna().sum()) for column in dataframe.columns],
            "Nilai unik": [int(dataframe[column].nunique(dropna=True)) for column in dataframe.columns],
        }
    )
    numeric_summary = (
        numeric_data.describe().transpose().reset_index(names="Kolom")
        if numeric_columns
        else pd.DataFrame()
    )

    categorical_frequency: dict[str, pd.DataFrame] = {}
    for column in categorical_columns:
        frequencies = dataframe[column].value_counts(dropna=False).head(10)
        categorical_frequency[str(column)] = frequencies.rename_axis(
            "Kategori"
        ).reset_index(name="Frekuensi")

    test_results: list[dict[str, object]] = []
    for column in numeric_columns[:20]:
        values = numeric_data[column].dropna()
        if len(values) < 3 or values.nunique() < 2:
            continue
        shapiro_values = (
            values.sample(n=5000, random_state=42)
            if len(values) > 5000
            else values
        )
        statistic, p_value = stats.shapiro(shapiro_values)
        test_results.append(
            {
                "Uji": "Shapiro-Wilk",
                "Variabel": str(column),
                "Perbandingan": "Normalitas",
                "N": len(shapiro_values),
                "Statistik": float(statistic),
                "p-value": float(p_value),
            }
        )

    correlation_pairs = itertools.islice(
        itertools.combinations(numeric_columns[:20], 2),
        30,
    )
    for first_column, second_column in correlation_pairs:
        paired_values = numeric_data[[first_column, second_column]].dropna()
        if len(paired_values) < 3:
            continue
        if (
            paired_values[first_column].nunique() < 2
            or paired_values[second_column].nunique() < 2
        ):
            continue
        statistic, p_value = stats.spearmanr(
            paired_values[first_column],
            paired_values[second_column],
        )
        test_results.append(
            {
                "Uji": "Korelasi Spearman",
                "Variabel": str(first_column),
                "Perbandingan": str(second_column),
                "N": len(paired_values),
                "Statistik": float(statistic),
                "p-value": float(p_value),
            }
        )

    tested_group_pairs = 0
    for category_column in categorical_columns:
        groups = dataframe[category_column].dropna().unique()
        if len(groups) < 2 or len(groups) > 10:
            continue
        for numeric_column in numeric_columns:
            grouped_values = [
                numeric_data.loc[dataframe[category_column] == group, numeric_column]
                .dropna()
                .to_numpy(dtype=float)
                for group in groups
            ]
            if any(len(group_values) < 2 for group_values in grouped_values):
                continue
            if np.unique(np.concatenate(grouped_values)).size < 2:
                continue

            if len(groups) == 2:
                test_name = "Mann-Whitney U"
                statistic, p_value = stats.mannwhitneyu(
                    grouped_values[0],
                    grouped_values[1],
                    alternative="two-sided",
                )
            else:
                test_name = "Kruskal-Wallis"
                statistic, p_value = stats.kruskal(*grouped_values)
            test_results.append(
                {
                    "Uji": test_name,
                    "Variabel": str(numeric_column),
                    "Perbandingan": str(category_column),
                    "N": sum(len(group_values) for group_values in grouped_values),
                    "Statistik": float(statistic),
                    "p-value": float(p_value),
                }
            )
            tested_group_pairs += 1
            if tested_group_pairs >= 30:
                break
        if tested_group_pairs >= 30:
            break

    tests_dataframe = pd.DataFrame(
        test_results,
        columns=["Uji", "Variabel", "Perbandingan", "N", "Statistik", "p-value"],
    )
    if not tests_dataframe.empty:
        tests_dataframe["Interpretasi awal"] = tests_dataframe["p-value"].apply(
            lambda value: (
                "p < 0,05; indikasi signifikan"
                if value < 0.05
                else "p >= 0,05; belum ada bukti signifikan"
            )
        )

    return {
        "column_profile": column_profile,
        "numeric_summary": numeric_summary,
        "categorical_frequency": categorical_frequency,
        "tests": tests_dataframe,
    }


def run_requested_statistical_test(
    dataframe: pd.DataFrame,
    analysis_plan: dict[str, object],
) -> dict[str, object]:
    test_name = analysis_plan.get("test")
    columns = analysis_plan.get("columns")
    group_column = analysis_plan.get("group_column")
    if not isinstance(test_name, str) or not isinstance(columns, list):
        raise ValueError("Rencana analisis AI tidak memiliki metode atau kolom valid.")
    if len(set(columns)) != len(columns):
        raise ValueError("Kolom yang sama tidak boleh dipilih lebih dari satu kali.")
    if any(not isinstance(column, str) or column not in dataframe.columns for column in columns):
        raise ValueError("Rencana analisis menyebut kolom yang tidak ada di CSV.")
    if (
        group_column is not None
        and (
            not isinstance(group_column, str)
            or group_column not in dataframe.columns
        )
    ):
        raise ValueError("Kolom kelompok yang dipilih AI tidak ada di CSV.")
    if group_column in columns:
        raise ValueError("Kolom kelompok harus berbeda dari kolom numerik yang diuji.")

    if test_name == "chi_square":
        if len(columns) != 2 or group_column is not None:
            raise ValueError("Uji Chi-Square memerlukan tepat dua kolom kategori.")
        contingency_table = pd.crosstab(
            dataframe[columns[0]],
            dataframe[columns[1]],
        )
        if min(contingency_table.shape) < 2:
            raise ValueError("Chi-Square memerlukan minimal dua kategori pada tiap kolom.")
        statistic, p_value, degrees_of_freedom, _ = stats.chi2_contingency(
            contingency_table
        )
        sample_size = int(contingency_table.to_numpy().sum())
        result_label = "Uji Chi-Square independensi"
        extra_result = f"Derajat kebebasan: {degrees_of_freedom}"
    elif test_name in {"pearson", "spearman"}:
        if len(columns) != 2 or group_column is not None:
            raise ValueError("Uji korelasi memerlukan tepat dua kolom numerik.")
        if any(not pd.api.types.is_numeric_dtype(dataframe[column]) for column in columns):
            raise ValueError("Uji korelasi hanya dapat digunakan pada kolom numerik.")
        paired_values = dataframe[columns].replace(
            [np.inf, -np.inf],
            np.nan,
        ).dropna()
        if len(paired_values) < 3:
            raise ValueError("Uji korelasi memerlukan minimal tiga pasangan data lengkap.")
        if any(paired_values[column].nunique() < 2 for column in columns):
            raise ValueError("Uji korelasi memerlukan variasi pada kedua kolom.")
        if test_name == "pearson":
            statistic, p_value = stats.pearsonr(
                paired_values[columns[0]],
                paired_values[columns[1]],
            )
            result_label = "Korelasi Pearson"
        else:
            statistic, p_value = stats.spearmanr(
                paired_values[columns[0]],
                paired_values[columns[1]],
            )
            result_label = "Korelasi Spearman"
        sample_size = len(paired_values)
        extra_result = ""
    elif test_name in {
        "independent_t",
        "mann_whitney",
        "one_way_anova",
        "kruskal",
    }:
        if len(columns) != 1 or not isinstance(group_column, str):
            raise ValueError(
                "Uji beda memerlukan satu kolom numerik dan satu kolom kelompok."
            )
        numeric_column = columns[0]
        if not pd.api.types.is_numeric_dtype(dataframe[numeric_column]):
            raise ValueError(f"Kolom `{numeric_column}` harus bertipe numerik.")
        grouped_data = dataframe[[numeric_column, group_column]].dropna()
        grouped_data[numeric_column] = grouped_data[numeric_column].replace(
            [np.inf, -np.inf],
            np.nan,
        )
        grouped_data = grouped_data.dropna()
        groups = grouped_data[group_column].unique()
        if len(groups) < 2 or len(groups) > 10:
            raise ValueError("Uji beda memerlukan 2 sampai 10 kelompok.")
        group_values = [
            grouped_data.loc[
                grouped_data[group_column] == group,
                numeric_column,
            ].to_numpy(dtype=float)
            for group in groups
        ]
        if any(len(values) < 2 for values in group_values):
            raise ValueError("Setiap kelompok harus memiliki minimal dua data.")
        if np.unique(np.concatenate(group_values)).size < 2:
            raise ValueError("Uji beda memerlukan variasi nilai numerik.")

        if test_name == "independent_t":
            if len(groups) != 2:
                raise ValueError("Uji t independen memerlukan tepat dua kelompok.")
            statistic, p_value = stats.ttest_ind(
                group_values[0],
                group_values[1],
                equal_var=False,
            )
            result_label = "Uji t independen Welch"
        elif test_name == "mann_whitney":
            if len(groups) != 2:
                raise ValueError("Mann-Whitney U memerlukan tepat dua kelompok.")
            statistic, p_value = stats.mannwhitneyu(
                group_values[0],
                group_values[1],
                alternative="two-sided",
            )
            result_label = "Mann-Whitney U"
        elif test_name == "one_way_anova":
            if len(groups) < 2:
                raise ValueError("ANOVA satu arah memerlukan minimal dua kelompok.")
            statistic, p_value = stats.f_oneway(*group_values)
            result_label = "ANOVA satu arah"
        else:
            statistic, p_value = stats.kruskal(*group_values)
            result_label = "Kruskal-Wallis"
        sample_size = len(grouped_data)
        extra_result = (
            "Kelompok: "
            + ", ".join(str(group) for group in groups)
        )
    else:
        raise ValueError(
            "Metode analisis tidak didukung. Metode yang diizinkan: "
            "Chi-Square, Pearson, Spearman, uji t Welch, Mann-Whitney U, "
            "ANOVA satu arah, dan Kruskal-Wallis."
        )

    if not math.isfinite(float(statistic)) or not math.isfinite(float(p_value)):
        raise ValueError("Uji menghasilkan nilai tidak valid untuk data yang dipilih.")
    return {
        "test": result_label,
        "columns": columns,
        "group_column": group_column,
        "n": sample_size,
        "statistic": float(statistic),
        "p_value": float(p_value),
        "extra": extra_result,
    }


def get_photo_history_directory(username: str) -> Path:
    username_hash = hashlib.sha256(username.encode("utf-8")).hexdigest()
    return PHOTO_HISTORY_DIR / username_hash


def load_photo_history(username: str) -> list[dict[str, str]]:
    photo_directory = get_photo_history_directory(username)
    history_path = photo_directory / "history.json"
    if not history_path.exists():
        return []

    document = json.loads(history_path.read_text(encoding="utf-8"))
    if not isinstance(document, dict) or not isinstance(document.get("images"), list):
        raise ValueError("Format riwayat foto tidak valid.")

    history = document["images"]
    for entry in history:
        if (
            not isinstance(entry, dict)
            or not isinstance(entry.get("id"), str)
            or not re.fullmatch(r"[a-f0-9]{32}", entry["id"])
            or not isinstance(entry.get("filename"), str)
            or not re.fullmatch(r"[a-f0-9]{32}\.(png|jpg|webp)", entry["filename"])
            or not isinstance(entry.get("prompt"), str)
            or not isinstance(entry.get("created_at"), str)
            or not isinstance(entry.get("mime_type"), str)
        ):
            raise ValueError("Entri pada riwayat foto tidak valid.")
    return history


def save_generated_photo(
    username: str,
    prompt: str,
    image_bytes: bytes,
    mime_type: str,
) -> dict[str, str]:
    extension_by_mime = {
        "image/png": "png",
        "image/jpeg": "jpg",
        "image/webp": "webp",
    }
    extension = extension_by_mime.get(mime_type)
    if extension is None:
        raise ValueError(f"Format foto keluaran tidak didukung: {mime_type}")

    photo_directory = get_photo_history_directory(username)
    photo_directory.mkdir(parents=True, exist_ok=True)
    image_id = uuid.uuid4().hex
    filename = f"{image_id}.{extension}"
    image_path = photo_directory / filename
    image_path.write_bytes(image_bytes)

    try:
        history = load_photo_history(username)
        entry = {
            "id": image_id,
            "filename": filename,
            "prompt": prompt,
            "created_at": datetime.now().astimezone().isoformat(timespec="seconds"),
            "mime_type": mime_type,
        }
        history.insert(0, entry)
        history_path = photo_directory / "history.json"
        temporary_path = photo_directory / f"{image_id}.tmp"
        temporary_path.write_text(
            json.dumps({"images": history}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        temporary_path.replace(history_path)
    except (OSError, json.JSONDecodeError, ValueError):
        image_path.unlink(missing_ok=True)
        raise
    return entry


def delete_generated_photo(username: str, image_id: str) -> None:
    if not re.fullmatch(r"[a-f0-9]{32}", image_id):
        raise ValueError("ID foto tidak valid.")

    photo_directory = get_photo_history_directory(username)
    history = load_photo_history(username)
    entry = next((item for item in history if item["id"] == image_id), None)
    if entry is None:
        raise FileNotFoundError("Foto tidak ditemukan pada riwayat pengguna.")

    image_path = photo_directory / entry["filename"]
    if image_path.exists():
        image_path.unlink()
    history = [item for item in history if item["id"] != image_id]
    history_path = photo_directory / "history.json"
    temporary_path = photo_directory / f"{image_id}.tmp"
    temporary_path.write_text(
        json.dumps({"images": history}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    temporary_path.replace(history_path)


def parse_visual_action(raw_response: str) -> tuple[str, dict[str, object] | None]:
    marker = "||VISUAL_ACTION||"
    if marker not in raw_response:
        return raw_response.strip(), None

    narrative, separator, payload = raw_response.partition(marker)
    if not separator:
        return raw_response.strip(), None
    json_payload, closing_marker, _ = payload.partition(marker)
    if not closing_marker:
        raise ValueError("Pembatas JSON visualisasi tidak lengkap.")

    try:
        instruction = json.loads(json_payload.strip())
    except json.JSONDecodeError:
        raise ValueError("JSON visualisasi dari AI tidak valid.") from None

    if not isinstance(instruction, dict):
        raise ValueError("Instruksi visualisasi harus berupa objek JSON.")
    action = instruction.get("action")
    if not isinstance(action, str) or action not in ("chart", "ppt", "both"):
        raise ValueError("Action visualisasi harus chart, ppt, atau both.")
    if action in ("chart", "both"):
        chart_type = instruction.get("chart_type")
        if not isinstance(chart_type, str) or chart_type not in ("bar", "pie"):
            raise ValueError("Jenis chart harus bar atau pie.")
    if action in {"ppt", "both"}:
        slides = instruction.get("ppt_slides")
        if not isinstance(slides, list) or not slides:
            raise ValueError("Instruksi PPT harus berisi minimal satu slide.")
        if len(slides) > 20:
            raise ValueError("Instruksi PPT melebihi batas 20 slide.")
        for slide in slides:
            if (
                not isinstance(slide, dict)
                or not isinstance(slide.get("judul", "Analisis Data"), str)
                or not isinstance(slide.get("poin_poin", []), list)
                or not all(isinstance(point, str) for point in slide.get("poin_poin", []))
            ):
                raise ValueError("Format salah satu slide PPT tidak valid.")
    return narrative.strip(), instruction


def remove_action_block(raw_response: str, marker: str) -> str:
    narrative, separator, payload = raw_response.partition(marker)
    if not separator:
        return raw_response.strip()
    _, closing_marker, suffix = payload.partition(marker)
    if not closing_marker:
        return narrative.strip()
    return f"{narrative}\n{suffix}".strip()


def parse_data_action(raw_response: str) -> tuple[str, dict[str, object] | None]:
    marker = "||DATA_ACTION||"
    if marker not in raw_response:
        return raw_response.strip(), None

    narrative, _, payload = raw_response.partition(marker)
    json_payload, closing_marker, suffix = payload.partition(marker)
    if not closing_marker:
        raise ValueError("Pembatas JSON aksi data tidak lengkap.")
    if marker in suffix:
        raise ValueError("Jawaban hanya boleh berisi satu blok aksi data.")
    try:
        instruction = json.loads(json_payload.strip())
    except json.JSONDecodeError:
        raise ValueError("JSON aksi data dari AI tidak valid.") from None
    if not isinstance(instruction, dict):
        raise ValueError("Instruksi data harus berupa objek JSON.")

    action = instruction.get("action")
    if not isinstance(action, str) or action not in {"table", "excel"}:
        raise ValueError("Aksi data harus berupa table atau excel.")
    title = instruction.get("title")
    if not isinstance(title, str) or not title.strip():
        raise ValueError("Judul tabel atau file Excel wajib diisi.")
    queries = instruction.get("queries")
    if not isinstance(queries, list) or not queries or len(queries) > 3:
        raise ValueError("Aksi data harus berisi 1 hingga 3 permintaan data.")
    seen_sources: set[tuple[str, str]] = set()
    for query in queries:
        if not isinstance(query, dict):
            raise ValueError("Format permintaan data tidak valid.")
        source = query.get("source")
        table = query.get("table", "")
        if source == "database":
            if not isinstance(table, str) or table not in {"survey_data", "raw_data"}:
                raise ValueError("Tabel database yang diminta tidak diizinkan.")
        elif source == "conversation":
            if table:
                raise ValueError("Permintaan conversation tidak boleh memilih tabel.")
        else:
            raise ValueError("Sumber data harus database atau conversation.")
        source_key = (source, table)
        if source_key in seen_sources:
            raise ValueError("Sumber data yang sama diminta lebih dari sekali.")
        seen_sources.add(source_key)
    return f"{narrative}\n{suffix}".strip(), instruction


def build_conversation_dataframe(
    messages: list[dict[str, object]],
) -> pd.DataFrame:
    rows = []
    for message_number, message in enumerate(messages, start=1):
        rows.append(
            {
                "message_number": message_number,
                "role": message.get("role", ""),
                "content": message.get("content", ""),
                "visual_action": json.dumps(
                    message.get("visual_action"),
                    ensure_ascii=False,
                )
                if message.get("visual_action") is not None
                else "",
                "data_action": json.dumps(
                    message.get("data_action"),
                    ensure_ascii=False,
                )
                if message.get("data_action") is not None
                else "",
            }
        )
    return pd.DataFrame(
        rows,
        columns=[
            "message_number",
            "role",
            "content",
            "visual_action",
            "data_action",
        ],
    )


def execute_data_query(
    query: dict[str, object],
    database_tables: dict[str, pd.DataFrame],
    conversation_dataframe: pd.DataFrame,
) -> pd.DataFrame:
    source = query.get("source")
    if source == "database":
        table = query.get("table")
        if not isinstance(table, str) or table not in {"survey_data", "raw_data"}:
            raise ValueError("Tabel database yang diminta tidak diizinkan.")
        source_dataframe = database_tables.get(table)
        if source_dataframe is None:
            raise ValueError(f"Tabel `{table}` tidak tersedia.")
    elif source == "conversation":
        source_dataframe = conversation_dataframe
    else:
        raise ValueError("Sumber data harus database atau conversation.")

    dataframe = source_dataframe.copy()
    filters = query.get("filters", [])
    if not isinstance(filters, list) or len(filters) > 20:
        raise ValueError("Filter harus berupa daftar maksimal 20 aturan.")
    for filter_rule in filters:
        if not isinstance(filter_rule, dict):
            raise ValueError("Format filter data tidak valid.")
        column = filter_rule.get("column")
        operator = filter_rule.get("operator")
        if not isinstance(operator, str):
            raise ValueError("Operator filter tidak valid.")
        if not isinstance(column, str) or column not in dataframe.columns:
            raise ValueError(f"Kolom filter tidak ditemukan: {column!r}.")
        series = dataframe[column]
        value = filter_rule.get("value")
        if operator == "equals":
            if not isinstance(value, (str, int, float, bool)):
                raise ValueError("Nilai filter equals tidak valid.")
            if pd.api.types.is_numeric_dtype(series) and isinstance(value, (int, float)):
                mask = series.eq(value)
            else:
                mask = series.astype("string").str.casefold().eq(str(value).casefold())
        elif operator == "not_equals":
            if not isinstance(value, (str, int, float, bool)):
                raise ValueError("Nilai filter not_equals tidak valid.")
            if pd.api.types.is_numeric_dtype(series) and isinstance(value, (int, float)):
                mask = series.ne(value)
            else:
                mask = series.astype("string").str.casefold().ne(str(value).casefold())
        elif operator == "contains":
            if not isinstance(value, str) or len(value) > 1000:
                raise ValueError("Nilai filter contains harus berupa teks.")
            mask = series.astype("string").str.contains(
                value,
                case=False,
                regex=False,
                na=False,
            )
        elif operator == "in":
            if (
                not isinstance(value, list)
                or not value
                or len(value) > 100
                or not all(isinstance(item, (str, int, float, bool)) for item in value)
            ):
                raise ValueError("Nilai filter in harus berisi 1 hingga 100 nilai.")
            if pd.api.types.is_numeric_dtype(series) and all(
                isinstance(item, (int, float)) and not isinstance(item, bool)
                for item in value
            ):
                mask = series.isin(value)
            else:
                accepted_values = {str(item).casefold() for item in value}
                mask = series.astype("string").str.casefold().isin(accepted_values)
        elif operator in {"gt", "gte", "lt", "lte"}:
            if not isinstance(value, (int, float)) or isinstance(value, bool):
                raise ValueError("Filter perbandingan hanya menerima nilai numerik.")
            numeric_series = pd.to_numeric(series, errors="coerce")
            nonempty_values = series.notna() & series.astype("string").str.strip().ne("")
            if (nonempty_values & numeric_series.isna()).any():
                raise ValueError(
                    f"Filter perbandingan hanya didukung untuk kolom numerik: {column}."
                )
            if operator == "gt":
                mask = numeric_series > value
            elif operator == "gte":
                mask = numeric_series >= value
            elif operator == "lt":
                mask = numeric_series < value
            else:
                mask = numeric_series <= value
        elif operator == "is_empty":
            mask = series.isna() | series.astype("string").str.strip().eq("")
        elif operator == "not_empty":
            mask = series.notna() & series.astype("string").str.strip().ne("")
        else:
            raise ValueError(f"Operator filter tidak didukung: {operator!r}.")
        dataframe = dataframe.loc[mask.fillna(False)]

    columns = query.get("columns", [])
    group_by = query.get("group_by", [])
    aggregations = query.get("aggregations", [])
    if not isinstance(columns, list) or len(columns) > 30:
        raise ValueError("Kolom yang diminta harus berupa daftar maksimal 30 kolom.")
    if not isinstance(group_by, list) or len(group_by) > 20:
        raise ValueError("Group_by harus berupa daftar maksimal 20 kolom.")
    if not isinstance(aggregations, list) or len(aggregations) > 20:
        raise ValueError("Agregasi harus berupa daftar maksimal 20 aturan.")
    for column in [*columns, *group_by]:
        if not isinstance(column, str) or column not in dataframe.columns:
            raise ValueError(f"Kolom data tidak ditemukan: {column!r}.")
    if len(set(group_by)) != len(group_by):
        raise ValueError("Kolom group_by tidak boleh berulang.")

    if group_by or aggregations:
        if group_by:
            grouped = dataframe.groupby(group_by, dropna=False, sort=False)
            if aggregations:
                result_parts = []
                used_output_names = set(group_by)
                for aggregation in aggregations:
                    if not isinstance(aggregation, dict):
                        raise ValueError("Format agregasi data tidak valid.")
                    column = aggregation.get("column")
                    function = aggregation.get("function")
                    alias = aggregation.get("alias")
                    if (
                        not isinstance(alias, str)
                        or not alias.strip()
                        or alias.strip() in used_output_names
                    ):
                        raise ValueError("Nama kolom hasil agregasi wajib diisi.")
                    used_output_names.add(alias.strip())
                    if function == "count" and column == "*":
                        result = grouped.size().rename(alias.strip())
                    else:
                        if not isinstance(column, str) or column not in dataframe.columns:
                            raise ValueError(f"Kolom agregasi tidak ditemukan: {column!r}.")
                        if not isinstance(function, str) or function not in {
                            "count",
                            "nunique",
                            "sum",
                            "mean",
                            "min",
                            "max",
                        }:
                            raise ValueError(f"Fungsi agregasi tidak didukung: {function!r}.")
                        if function in {"sum", "mean"} and not pd.api.types.is_numeric_dtype(
                            dataframe[column]
                        ):
                            raise ValueError(
                                f"Agregasi {function} hanya didukung untuk kolom numerik: {column}."
                            )
                        result = grouped[column].agg(function).rename(alias.strip())
                    result_parts.append(result)
                dataframe = pd.concat(result_parts, axis=1).reset_index()
            else:
                dataframe = grouped.size().rename("Jumlah").reset_index()
        else:
            result_values: dict[str, object] = {}
            for aggregation in aggregations:
                if not isinstance(aggregation, dict):
                    raise ValueError("Format agregasi data tidak valid.")
                column = aggregation.get("column")
                function = aggregation.get("function")
                alias = aggregation.get("alias")
                if (
                    not isinstance(alias, str)
                    or not alias.strip()
                    or alias.strip() in result_values
                ):
                    raise ValueError("Nama kolom hasil agregasi wajib diisi.")
                if function == "count" and column == "*":
                    result_values[alias.strip()] = len(dataframe)
                    continue
                if not isinstance(column, str) or column not in dataframe.columns:
                    raise ValueError(f"Kolom agregasi tidak ditemukan: {column!r}.")
                if not isinstance(function, str) or function not in {
                    "count",
                    "nunique",
                    "sum",
                    "mean",
                    "min",
                    "max",
                }:
                    raise ValueError(f"Fungsi agregasi tidak didukung: {function!r}.")
                if function in {"sum", "mean"} and not pd.api.types.is_numeric_dtype(
                    dataframe[column]
                ):
                    raise ValueError(
                        f"Agregasi {function} hanya didukung untuk kolom numerik: {column}."
                    )
                result_values[alias.strip()] = getattr(dataframe[column], function)()
            dataframe = pd.DataFrame([result_values])
    else:
        if not columns:
            raise ValueError("Pilih minimal satu kolom untuk menampilkan data.")
        dataframe = dataframe.loc[:, columns]

    sort_by = query.get("sort_by", [])
    if not isinstance(sort_by, list) or len(sort_by) > 10:
        raise ValueError("Pengurutan harus berupa daftar maksimal 10 aturan.")
    sort_columns = []
    ascending = []
    for sort_rule in sort_by:
        if not isinstance(sort_rule, dict):
            raise ValueError("Format aturan pengurutan tidak valid.")
        column = sort_rule.get("column")
        direction = sort_rule.get("direction", "asc")
        if not isinstance(column, str) or column not in dataframe.columns:
            raise ValueError(f"Kolom pengurutan tidak ditemukan: {column!r}.")
        if not isinstance(direction, str) or direction not in {"asc", "desc"}:
            raise ValueError("Arah pengurutan harus asc atau desc.")
        sort_columns.append(column)
        ascending.append(direction == "asc")
    if sort_columns:
        dataframe = dataframe.sort_values(sort_columns, ascending=ascending)

    limit = query.get("limit", 1000)
    if not isinstance(limit, int) or isinstance(limit, bool) or not 1 <= limit <= 5000:
        raise ValueError("Batas hasil harus bilangan bulat antara 1 hingga 5000.")
    return dataframe.head(limit).reset_index(drop=True)


def render_data_action(
    instruction: dict[str, object],
    database_tables: dict[str, pd.DataFrame],
    conversation_dataframe: pd.DataFrame,
    message_index: int,
) -> None:
    action = instruction.get("action")
    queries = instruction.get("queries")
    title = instruction.get("title")
    if (
        not isinstance(action, str)
        or action not in {"table", "excel"}
        or not isinstance(title, str)
        or not isinstance(queries, list)
        or not queries
        or len(queries) > 3
    ):
        st.warning("Instruksi tabel atau Excel pada riwayat chat tidak valid.")
        return

    results: list[tuple[str, pd.DataFrame]] = []
    try:
        for query in queries:
            if not isinstance(query, dict):
                raise ValueError("Format permintaan data tidak valid.")
            source_label = (
                f"Database: {query.get('table')}"
                if query.get("source") == "database"
                else "Riwayat percakapan"
            )
            results.append(
                (
                    source_label,
                    execute_data_query(
                        query,
                        database_tables,
                        conversation_dataframe,
                    ),
                )
            )
    except ValueError as error:
        st.warning(f"Tidak dapat menampilkan permintaan data: {error}")
        return

    for result_index, (source_label, dataframe) in enumerate(results):
        st.markdown(f"**{source_label}**")
        if dataframe.empty:
            st.info("Tidak ada data yang cocok dengan permintaan.")
        else:
            st.dataframe(dataframe, alt=f"Hasil {source_label.lower()}")

    if action == "excel":
        safe_title = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", title).strip(" .")
        if not safe_title:
            safe_title = "Hasil_Tabulasi"
        excel_buffer = io.BytesIO()
        with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
            used_sheet_names: set[str] = set()
            for result_index, (source_label, dataframe) in enumerate(results, start=1):
                base_sheet_name = (
                    "Percakapan"
                    if source_label == "Riwayat percakapan"
                    else str(source_label.split(": ", maxsplit=1)[-1])
                )
                base_sheet_name = re.sub(r"[\\/*?:\[\]]", "_", base_sheet_name)[:31]
                sheet_name = base_sheet_name or f"Data_{result_index}"
                unique_sheet_name = sheet_name
                suffix = 2
                while unique_sheet_name.casefold() in used_sheet_names:
                    suffix_text = f"_{suffix}"
                    unique_sheet_name = f"{sheet_name[:31-len(suffix_text)]}{suffix_text}"
                    suffix += 1
                used_sheet_names.add(unique_sheet_name.casefold())
                excel_dataframe = dataframe.copy()
                for column in excel_dataframe.select_dtypes(
                    include=["object", "string"]
                ).columns:
                    excel_dataframe[column] = excel_dataframe[column].map(
                        lambda value: (
                            f"'{value}"
                            if isinstance(value, str)
                            and value.startswith(("=", "+", "-", "@"))
                            else value
                        )
                    )
                excel_dataframe.to_excel(
                    writer,
                    sheet_name=unique_sheet_name,
                    index=False,
                )
        st.download_button(
            label=f"Unduh {safe_title}.xlsx",
            data=excel_buffer.getvalue(),
            file_name=f"{safe_title}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key=f"data_excel_download_{message_index}",
        )


def render_visual_action(
    instruction: dict[str, object],
    answer_counts: dict[str, int],
    message_index: int,
) -> None:
    action = instruction.get("action")
    if not isinstance(action, str) or action not in ("chart", "ppt", "both"):
        st.warning("Instruksi visualisasi pada riwayat chat tidak valid.")
        return
    title = instruction.get("title")
    if not isinstance(title, str) or not title.strip():
        title = "Analisis Data"
    safe_title = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", title).strip(" .")
    if not safe_title:
        safe_title = "Analisis_Data"

    if action in {"chart", "both"}:
        if not answer_counts:
            st.info("Tidak ada data jawaban untuk dibuatkan grafik.")
        else:
            chart_type = instruction.get("chart_type")
            chart_data = pd.DataFrame(
                list(answer_counts.items()),
                columns=["Jawaban", "Jumlah"],
            )
            figure, axis = plt.subplots(figsize=(8, 4))
            if chart_type == "pie":
                axis.pie(
                    chart_data["Jumlah"],
                    labels=chart_data["Jawaban"],
                    autopct="%1.1f%%",
                    startangle=90,
                    colors=sns.color_palette("pastel"),
                )
                axis.axis("equal")
            else:
                sns.barplot(
                    x="Jumlah",
                    y="Jawaban",
                    data=chart_data,
                    hue="Jawaban",
                    palette="viridis",
                    legend=False,
                    ax=axis,
                )
            axis.set_title(title)
            figure.tight_layout()
            st.pyplot(figure)
            plt.close(figure)

    if action in {"ppt", "both"}:
        slides = instruction.get("ppt_slides")
        if (
            not isinstance(slides, list)
            or not slides
            or len(slides) > 20
            or any(
                not isinstance(slide, dict)
                or not isinstance(slide.get("judul", "Analisis Data"), str)
                or not isinstance(slide.get("poin_poin", []), list)
                or not all(
                    isinstance(point, str)
                    for point in slide.get("poin_poin", [])
                )
                for slide in slides
            )
        ):
            st.info("Instruksi PPT tidak memiliki slide yang valid untuk dibuat.")
            return

        from pptx import Presentation
        from pptx.util import Inches, Pt

        presentation = Presentation()
        for slide_data in slides:
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            title_box = slide.shapes.add_textbox(
                Inches(0.5),
                Inches(0.5),
                Inches(9),
                Inches(1),
            )
            title_paragraph = title_box.text_frame.paragraphs[0]
            title_paragraph.text = slide_data.get("judul", "Analisis Data")
            title_paragraph.font.size = Pt(28)

            body_box = slide.shapes.add_textbox(
                Inches(0.75),
                Inches(1.8),
                Inches(8.5),
                Inches(5),
            )
            body_frame = body_box.text_frame
            body_frame.word_wrap = True
            for point_index, point in enumerate(slide_data.get("poin_poin", [])):
                paragraph = (
                    body_frame.paragraphs[0]
                    if point_index == 0
                    else body_frame.add_paragraph()
                )
                paragraph.text = point
                paragraph.level = 0
                paragraph.font.size = Pt(20)

        ppt_buffer = io.BytesIO()
        presentation.save(ppt_buffer)
        st.download_button(
            label=f"Download {safe_title}.pptx",
            data=ppt_buffer.getvalue(),
            file_name=f"{safe_title}.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            key=f"ppt_download_{message_index}",
        )


# --- 1. INISIALISASI MEMORI (SESSION STATE) ---
initialize_user_database()

if "authenticated_username" not in st.session_state:
    st.session_state.authenticated_username = None
if "authenticated_role" not in st.session_state:
    st.session_state.authenticated_role = None

if st.session_state.authenticated_role is None:
    _, center, _ = st.columns([1, 1.2, 1]) 
    with center: 
        st.markdown(
            '<div class="login-brand"><h1>AI Data Analyst Workspace</h1>'
            '<p>Masuk untuk mulai menganalisis data survei</p></div>',
            unsafe_allow_html=True,
        )
        with st.form("login_form"):
            login_username = st.text_input("Username")
            login_password = st.text_input("Password", type="password")
            login_submitted = st.form_submit_button(
                "Masuk", type="primary", width="stretch"
            )

    if login_submitted:
        role = authenticate_user(login_username.strip().lower(), login_password)
        if role is None:
            st.error("Username atau password salah.")
        else:
            st.session_state.authenticated_username = login_username.strip().lower()
            st.session_state.authenticated_role = role
            st.rerun()

    st.stop()

role = st.session_state.authenticated_role
navigation_pages = []
if role == "admin":
    navigation_pages.append(("Kelola User", "users"))
if role in {"admin", "statistician"}:
    navigation_pages.append(("Upload Data Project", "upload"))
navigation_pages.append(("Analisis AI", "analysis"))
navigation_pages.append(("Create / Edit Foto", "photo"))
navigation_pages.append(("Analisis statistik", "statistics"))

navigation_labels = [label for label, _ in navigation_pages]
if st.session_state.get("app_navigation") not in navigation_labels:
    st.session_state.app_navigation = navigation_labels[0]

with st.sidebar:
    with st.container(horizontal_alignment="center"):
        st.image(str(LOGO_PATH), width=138)
    st.caption(f"{st.session_state.authenticated_username} · {role}")
    st.divider()
    for page_label, page_key in navigation_pages:
        if st.button(
            page_label,
            key=f"navigation_{page_key}",
            type=(
                "primary"
                if st.session_state.app_navigation == page_label
                else "secondary"
            ),
            width="stretch",
        ):
            st.session_state.app_navigation = page_label
            st.rerun()

    selected_page_label = st.session_state.app_navigation
    selected_page = dict(navigation_pages)[selected_page_label]
    st.divider()
    if st.button("Keluar", icon=":material/logout:", width="stretch"):
        st.session_state.authenticated_username = None
        st.session_state.authenticated_role = None
        st.rerun()
    st.markdown(
        '<div class="sidebar-app-brand">AI Data Analyst<br>Workspace</div>',
        unsafe_allow_html=True,
    )

page_descriptions = {
    "users": "Kelola akun pengguna dan hak aksesnya.",
    "upload": (
        "Unggah Excel dengan baris 1 sebagai nama kolom, baris 2 sebagai "
        "deskripsi kolom, baris 3 sebagai format jawaban, dan baris 4 hingga "
        "akhir sebagai data."
    ),
    "analysis": "Analisis data survei bersama AI chat.",
    "statistics": (
        "Unggah CSV untuk analisis statistik otomatis dan tanya jawab dengan Gemini AI."
    ),
    "photo": "Buat atau edit gambar, lalu kelola riwayat foto Anda.",
}
st.title(selected_page_label)
st.caption(page_descriptions[selected_page])

if selected_page == "users":
    if role != "admin":
        st.error("Hanya admin yang dapat mengelola pengguna.")
        st.stop()

    with st.form("user_registration"):
        username = st.text_input("Username")
        password = st.text_input("Password", type="password")
        user_role = st.selectbox("User type / role", USER_ROLES)
        submitted = st.form_submit_button("Submit", type="primary")

    if submitted:
        username = username.strip().lower()
        if not username or not password:
            st.error("Username dan password wajib diisi.")
        else:
            try:
                save_user(username, password, user_role)
            except sqlite3.IntegrityError:
                st.error("Username tersebut sudah terdaftar.")
            except ValueError as error:
                st.error(str(error))
            except sqlite3.Error:
                st.error("Gagal menyimpan pengguna ke user.db. Periksa akses folder aplikasi.")
            else:
                st.success(f"Username berhasil disimpan dengan role {user_role}.")

    st.subheader("Daftar pengguna")
    try:
        users_dataframe = list_users()
    except (OSError, sqlite3.Error, pd.errors.DatabaseError) as error:
        st.error(f"Gagal membaca daftar pengguna dari user.db: {error}")
    else:
        if users_dataframe.empty:
            st.info("Belum ada pengguna terdaftar.")
        else:
            st.dataframe(
                users_dataframe,
                hide_index=True,
                alt="Daftar username dan role yang tersimpan di user.db",
            )

            usernames = users_dataframe["username"].tolist()
            selected_username = st.selectbox(
                "Pilih pengguna untuk dikelola",
                options=usernames,
                key="managed_username",
            )
            current_role = users_dataframe.loc[
                users_dataframe["username"] == selected_username,
                "role",
            ].iloc[0]
            role_widget_key = (
                "managed_role_"
                + hashlib.sha256(selected_username.encode("utf-8")).hexdigest()[:12]
            )

            with st.form("manage_user_form"):
                new_role = st.selectbox(
                    "Role pengguna",
                    options=USER_ROLES,
                    index=USER_ROLES.index(current_role),
                    key=role_widget_key,
                )
                confirm_delete = st.checkbox(
                    f"Saya yakin ingin menghapus akun {selected_username}."
                )
                save_role_submitted = st.form_submit_button(
                    "Simpan perubahan role",
                    type="primary",
                )
                delete_user_submitted = st.form_submit_button(
                    "Hapus username",
                    icon=":material/delete:",
                )

            if save_role_submitted:
                try:
                    update_user_role(
                        selected_username,
                        new_role,
                        st.session_state.authenticated_username,
                    )
                except (FileNotFoundError, PermissionError, ValueError) as error:
                    st.error(f"Gagal mengubah role: {error}")
                except sqlite3.Error as error:
                    st.error(f"Gagal menyimpan role ke user.db: {error}")
                else:
                    st.success(f"Role {selected_username} berhasil diubah menjadi {new_role}.")
                    st.rerun()

            if delete_user_submitted:
                if not confirm_delete:
                    st.error("Centang konfirmasi sebelum menghapus username.")
                else:
                    try:
                        delete_user(
                            selected_username,
                            st.session_state.authenticated_username,
                        )
                    except (FileNotFoundError, PermissionError, ValueError) as error:
                        st.error(f"Gagal menghapus pengguna: {error}")
                    except sqlite3.Error as error:
                        st.error(f"Gagal menghapus pengguna dari user.db: {error}")
                    else:
                        st.success(f"Akun {selected_username} berhasil dihapus.")
                        st.rerun()

# ==============================================================================
# TAB 2: UPLOAD DAN TRANSFORMASI DATA
# ==============================================================================
if selected_page == "upload":
        uploaded_file = st.file_uploader(
            "Unggah file Excel",
            type=["xlsx", "xls"],
        )

        if uploaded_file is not None:
            if st.button("Proses & Transformasi ke Vertikal 🚀", type="primary"):
                with st.spinner("Sedang merapikan data..."):
                    uploaded_file.seek(0)
                    excel_sheet = pd.read_excel(
                        uploaded_file,
                        header=None,
                        dtype=object,
                    )
                    if excel_sheet.shape[0] < 4 or excel_sheet.shape[1] < 3:
                        st.error(
                            "File Excel harus berisi minimal empat baris dan tiga kolom: "
                            "baris 1 nama kolom, baris 2 deskripsi kolom, baris 3 "
                            "format jawaban, dan baris 4 atau seterusnya data."
                        )
                    else:
                        headers = []
                        for column_number, value in enumerate(excel_sheet.iloc[0], start=1):
                            if pd.isna(value) or not str(value).strip():
                                headers.append(f"Kolom_{column_number}")
                            else:
                                headers.append(str(value).strip())

                        if len({header.casefold() for header in headers}) != len(headers):
                            st.error(
                                "Nama kolom pada baris pertama harus unik, termasuk "
                                "perbedaan huruf besar/kecil."
                            )
                            st.stop()

                        column_descriptions = [
                            "" if pd.isna(value) else str(value).strip()
                            for value in excel_sheet.iloc[1]
                        ]
                        answer_formats = [
                            "" if pd.isna(value) else str(value).strip()
                            for value in excel_sheet.iloc[2]
                        ]
                        description_dataframe = pd.DataFrame({
                            "no_kues": range(1, len(headers) + 1),
                            "nama_kolom": headers,
                            "deskripsi_kolom": column_descriptions,
                            "format_jawaban": answer_formats,
                        })
                        df_raw = excel_sheet.iloc[3:].copy()
                        df_raw.columns = headers
                        df_raw.reset_index(drop=True, inplace=True)
                        answer_formats_by_column = dict(zip(headers, answer_formats))

                        list_hasil_vertikal = []
                        kolom_id = df_raw.columns[0]
                        kolom_produk = df_raw.columns[1]
                        reserved_columns = {
                            "id",
                            "produk",
                            "jawaban",
                            "kolom jawaban",
                        }
                        source_column_names = {}
                        used_column_names = set()
                        for column in df_raw.columns:
                            output_name = column
                            if output_name.casefold() in reserved_columns:
                                output_name = f"{output_name}_sumber"
                            while (
                                output_name.casefold() in used_column_names
                                or output_name.casefold() in reserved_columns
                            ):
                                output_name = f"_{output_name}"
                            source_column_names[column] = output_name
                            used_column_names.add(output_name.casefold())

                        kolom_pertanyaan_normal = "Kolom Jawaban"
                        while (
                            kolom_pertanyaan_normal.casefold() in used_column_names
                        ):
                            kolom_pertanyaan_normal = f"_{kolom_pertanyaan_normal}"

                        for _, row in df_raw.iterrows():
                            for kolom_jawaban in df_raw.columns[2:]:
                                nilai = row[kolom_jawaban]
                                if pd.isna(nilai) or not str(nilai).strip():
                                    continue

                                format_jawaban = answer_formats_by_column[
                                    kolom_jawaban
                                ].casefold()
                                if (
                                    "multiple_answer" in format_jawaban
                                    or "multiple answer" in format_jawaban
                                ):
                                    jawaban_list = str(nilai).split(";")
                                else:
                                    jawaban_list = [str(nilai)]

                                for jawaban in jawaban_list:
                                    if not jawaban.strip():
                                        continue
                                    record = {
                                        source_column_names[column]: row[column]
                                        for column in df_raw.columns
                                    }
                                    for kolom_lain in df_raw.columns[2:]:
                                        if kolom_lain != kolom_jawaban:
                                            record[source_column_names[kolom_lain]] = None
                                    jawaban_bersih = jawaban.strip()
                                    record[source_column_names[kolom_jawaban]] = jawaban_bersih
                                    record["ID"] = row[kolom_id]
                                    record["Produk"] = row[kolom_produk]
                                    record["Jawaban"] = jawaban_bersih
                                    record[kolom_pertanyaan_normal] = kolom_jawaban
                                    list_hasil_vertikal.append(record)

                        output_columns = list(source_column_names.values())
                        for column in ("ID", "Produk", "Jawaban", kolom_pertanyaan_normal):
                            if column not in output_columns:
                                output_columns.append(column)
                        df_vertikal = pd.DataFrame(
                            list_hasil_vertikal,
                            columns=output_columns,
                        )
                        try:
                            database_path, description_path = save_project_database(
                                df_vertikal,
                                uploaded_file.name,
                                raw_dataframe=df_raw,
                                description_dataframe=description_dataframe,
                            )
                        except (sqlite3.Error, ValueError, pd.errors.DatabaseError) as error:
                            st.error(f"Gagal menyimpan database proyek: {error}")
                        else:
                            st.session_state.last_processed_db = database_path.name
                            st.session_state.selected_project_db = database_path.name
                            st.success(
                                "🎉 Dua database berhasil disimpan: "
                                f"`data_project\\{database_path.name}` dan "
                                f"`data_project\\{description_path.name}`. "
                                "Jika nama sama sudah ada, database sebelumnya diganti."
                            )
                
        # Tampilkan pratinjau data jika sudah diproses
        last_processed_db = st.session_state.get("last_processed_db")
        available_databases = {
            database.name: database for database in list_project_databases()
        }
        if last_processed_db in available_databases:
            st.subheader("Pratinjau Data Vertikal yang Siap Dianalisis:")
            try:
                df_preview = load_project_database(
                    available_databases[last_processed_db]
                )
            except (sqlite3.Error, pd.errors.DatabaseError) as error:
                st.error(f"Gagal membaca database proyek: {error}")
            else:
                st.dataframe(df_preview)

        st.divider()
        st.subheader("Database yang tersedia")
        project_databases = list_project_databases()
        if not project_databases:
            st.info("Belum ada database di folder data_project.")
        else:
            database_rows = []
            for database_path in project_databases:
                try:
                    submitted_at, updated_at = get_project_database_timestamps(
                        database_path
                    )
                except (OSError, sqlite3.Error, ValueError) as error:
                    st.error(
                        f"Gagal membaca tanggal database "
                        f"`{database_path.name}`: {error}"
                    )
                    continue
                database_rows.append(
                    {
                        "Database": database_path.name,
                        "Tanggal submit": submitted_at.strftime(
                            "%d-%m-%Y %H:%M:%S"
                        ),
                        "Tanggal update": updated_at.strftime(
                            "%d-%m-%Y %H:%M:%S"
                        ),
                    }
                )

            if database_rows:
                st.dataframe(
                    pd.DataFrame(database_rows),
                    hide_index=True,
                    alt="Database proyek beserta tanggal submit dan update",
                )

                database_names = [
                    database_path.name for database_path in project_databases
                ]
                database_to_delete = st.selectbox(
                    "Pilih database yang akan dihapus",
                    options=database_names,
                    key="project_database_to_delete",
                )
                confirmation_key = (
                    "confirm_delete_project_"
                    + hashlib.sha256(database_to_delete.encode("utf-8")).hexdigest()[:12]
                )
                confirm_database_delete = st.checkbox(
                    f"Saya yakin ingin menghapus `{database_to_delete}` "
                    "beserta file metadata-nya.",
                    key=confirmation_key,
                )
                if st.button(
                    "Hapus database",
                    icon=":material/delete:",
                    key="delete_project_database",
                    disabled=not confirm_database_delete,
                ):
                    try:
                        delete_project_database(database_to_delete)
                    except (FileNotFoundError, OSError, ValueError, sqlite3.Error) as error:
                        st.error(
                            f"Gagal menghapus database "
                            f"`{database_to_delete}`: {error}"
                        )
                    else:
                        if (
                            st.session_state.get("selected_project_db")
                            == database_to_delete
                        ):
                            st.session_state.selected_project_db = None
                        if (
                            st.session_state.get("last_processed_db")
                            == database_to_delete
                        ):
                            st.session_state.last_processed_db = None
                        st.success(
                            f"Database `{database_to_delete}` dan file "
                            "metadata-nya berhasil dihapus."
                        )
                        st.rerun()

# ==============================================================================
# TAB 3: TANYA AI, ANALISIS, DAN CHART
# ==============================================================================
if selected_page == "analysis":
    project_databases = list_project_databases()
    if not project_databases:
        st.warning("Belum ada database proyek. Proses data terlebih dahulu di Tab 2.")
    else:
        database_names = [database.name for database in project_databases]
        if st.session_state.get("selected_project_db") not in database_names:
            st.session_state.selected_project_db = database_names[0]

        database_search = st.text_input(
            "Cari database proyek",
            placeholder="Ketik nama database untuk mencari...",
            key="project_database_search",
        )
        matching_database_names = [
            database_name
            for database_name in database_names
            if database_search.casefold() in database_name.casefold()
        ]
        if database_search:
            if matching_database_names:
                if (
                    st.session_state.get("database_search_result")
                    not in matching_database_names
                ):
                    st.session_state.database_search_result = None
                search_result = st.selectbox(
                    "Hasil pencarian",
                    options=matching_database_names,
                    index=None,
                    placeholder="Pilih database dari hasil pencarian",
                    key="database_search_result",
                )
                if search_result is not None:
                    st.session_state.selected_project_db = search_result
            else:
                st.info(f"Tidak ada database yang cocok dengan “{database_search}”.")

        selected_database = st.selectbox(
            "Pilih database proyek",
            options=database_names,
            key="selected_project_db",
        )
        selected_database_path = DATA_PROJECT_DIR / selected_database

        try:
            df_selected = load_project_database(selected_database_path)
        except (sqlite3.Error, pd.errors.DatabaseError) as error:
            st.error(f"Gagal membaca database `{selected_database}`: {error}")
        else:
            database_tables = {"survey_data": df_selected}
            try:
                database_tables["raw_data"] = load_project_raw_database(
                    selected_database_path
                )
            except (sqlite3.Error, pd.errors.DatabaseError) as error:
                st.warning(f"Tabel raw_data tidak tersedia: {error}")

            tren_jawaban = (
                df_selected["Jawaban"].value_counts().head(50).to_dict()
                if "Jawaban" in df_selected.columns
                else {}
            )
            description_path = selected_database_path.with_name(
                f"{selected_database_path.stem}_desc.db"
            )
            try:
                question_descriptions = load_project_descriptions(
                    selected_database_path
                )
            except (sqlite3.Error, pd.errors.DatabaseError) as error:
                question_descriptions = pd.DataFrame(
                    columns=[
                        "no_kues",
                        "nama_kolom",
                        "deskripsi_kolom",
                        "format_jawaban",
                    ]
                )
                st.warning(
                    f"Metadata format jawaban dari `{description_path.name}` "
                    f"tidak tersedia atau gagal dibaca: {error}"
                )

            if question_descriptions.empty:
                description_context = (
                    "Metadata format jawaban tidak tersedia. Jangan menebak "
                    "apakah sebuah kolom single-answer atau multiple-answer."
                )
            else:
                description_lines = [
                    f"- no_kues {row.no_kues}, nama_kolom '{row.nama_kolom}', "
                    f"deskripsi '{row.deskripsi_kolom or 'tidak ada'}', "
                    f"format_jawaban "
                    f"'{row.format_jawaban or 'format tidak ditentukan'}'"
                    for row in question_descriptions.itertuples(index=False)
                ]
                description_context = (
                    f"Format jawaban setiap kolom wajib mengacu pada database "
                    f"metadata `{description_path.name}`. Hubungkan metadata "
                    "ke data menggunakan `nama_kolom`. Baris metadata berisi "
                    "`no_kues`, `nama_kolom` dari baris 1, `deskripsi_kolom` "
                    "dari baris 2, dan `format_jawaban` dari baris 3 Excel:\n"
                    + "\n".join(description_lines)
                    + "\nIkuti metadata tersebut: single_answer berarti satu "
                    "jawaban per responden; multiple_answer berarti beberapa "
                    "jawaban yang dipisahkan dengan titik koma (;)."
                )

            st.write(f"**Database aktif:** `{selected_database}`")
            try:
                chat_histories = load_chat_histories(
                    st.session_state.authenticated_username
                )
            except (OSError, json.JSONDecodeError, ValueError) as error:
                st.error(f"Gagal membaca riwayat chat: {error}")
            else:
                messages = chat_histories.get(selected_database, [])
                chat_container = st.container()
                prompt = st.chat_input("Tulis pesan tentang data ini...")
                chat_error = None

                if prompt:
                    messages.append({"role": "user", "content": prompt})
                    chat_histories[selected_database] = messages
                    try:
                        save_chat_histories(
                            st.session_state.authenticated_username,
                            chat_histories,
                        )
                    except OSError as error:
                        chat_error = f"Gagal menyimpan pesan ke riwayat chat: {error}"
                    else:
                        if "GEMINI_API_KEY" not in st.secrets:
                            chat_error = "API Key belum diset di Secrets Streamlit."
                        else:
                            conversation_json = json.dumps(
                                messages[:-1],
                                ensure_ascii=False,
                                default=str,
                            )
                            table_schema_lines = []
                            for table_name, table_dataframe in database_tables.items():
                                columns = ", ".join(
                                    f"{column} ({table_dataframe[column].dtype})"
                                    for column in table_dataframe.columns
                                )
                                table_description = (
                                    "data jawaban vertikal; satu responden dapat "
                                    "muncul di beberapa baris"
                                    if table_name == "survey_data"
                                    else "data Excel asli; satu baris per responden"
                                )
                                table_schema_lines.append(
                                    f"- {table_name}: {len(table_dataframe)} baris; "
                                    f"{table_description}; kolom: {columns}"
                                )
                            database_schema_context = "\n".join(table_schema_lines)
                            if tren_jawaban:
                                answer_distribution_context = json.dumps(
                                    tren_jawaban,
                                    ensure_ascii=False,
                                    default=str,
                                )
                            else:
                                answer_distribution_context = (
                                    "Tidak ada kolom bernama Jawaban. Gunakan "
                                    "kolom yang tercantum pada skema database dan "
                                    "aksi data terstruktur untuk menghitung tabulasi."
                                )
                            ai_prompt = f"""
                            Anda adalah Data Analyst profesional. Jawab pertanyaan
                            terbaru berdasarkan database proyek yang dipilih dan
                            riwayat percakapan JSON berikut. Jangan mengarang nilai
                            atau hasil query yang tidak tersedia.

                            Database aktif: {selected_database}
                            Skema dan jumlah baris database:
                            {database_schema_context}

                            Ringkasan distribusi jawaban yang tersedia:
                            {answer_distribution_context}

                            Panduan format jawaban kolom:
                            {description_context}

                            Riwayat percakapan sebelumnya dalam format JSON:
                            {conversation_json or '(Belum ada percakapan sebelumnya)'}

                            Pertanyaan terbaru: {prompt}

                            ATURAN VISUALISASI:
                            Jika pengguna meminta grafik/chart atau PowerPoint (PPT),
                            sertakan objek JSON mentah di bagian paling bawah jawaban
                            dengan format berikut. Ganti nilainya sesuai permintaan:
                            ||VISUAL_ACTION||
                            {{
                              "action": "chart",
                              "chart_type": "bar",
                              "title": "Judul grafik atau presentasi",
                              "ppt_slides": [
                                {{"judul": "Slide 1", "poin_poin": ["Insight A", "Insight B"]}}
                              ]
                            }}
                            ||VISUAL_ACTION||
                            Nilai action harus "chart", "ppt", atau "both".
                            chart_type harus "bar" atau "pie".
                            Untuk PPT, sertakan daftar ppt_slides dengan judul dan poin_poin.
                            JSON harus valid dan diapit pembatas tersebut. Jika pengguna
                            tidak meminta chart atau PPT, jangan sertakan pembatas atau JSON.

                            ATURAN TABULASI, PENGAMBILAN DATA, DAN EXCEL:
                            Jika pengguna meminta tabulasi, tabel data, daftar/baris
                            tertentu, atau ekstraksi data dari database/riwayat JSON,
                            sertakan satu blok JSON DATA_ACTION mentah di paling bawah
                            jawaban. Aksi ini mengambil data langsung dari sumber lokal;
                            jangan menulis hasil angka seolah sudah dihitung jika belum
                            tersedia pada ringkasan.
                            Format:
                            ||DATA_ACTION||
                            {{
                              "action": "table",
                              "title": "Judul hasil",
                              "queries": [
                                {{
                                  "source": "database",
                                  "table": "survey_data",
                                  "columns": ["Provinsi"],
                                  "filters": [],
                                  "group_by": ["Provinsi"],
                                  "aggregations": [
                                    {{"column": "*", "function": "count", "alias": "Jumlah"}}
                                  ],
                                  "sort_by": [{{"column": "Jumlah", "direction": "desc"}}],
                                  "limit": 1000
                                }}
                              ]
                            }}
                            ||DATA_ACTION||
                            Isi queries dengan satu hingga tiga permintaan. Untuk sumber
                            database, source hanya "database", table harus "survey_data"
                            atau "raw_data", dan nama kolom harus persis sesuai skema.
                            Untuk sumber percakapan JSON, gunakan source "conversation";
                            kolom yang tersedia: message_number, role, content,
                            visual_action, data_action. Minta kolom relevan saja.
                            Setiap query dapat memakai filters dengan operator equals,
                            not_equals, contains, in, gt, gte, lt, lte, is_empty, atau
                            not_empty. Tabulasi memakai group_by dan aggregations;
                            fungsi agregasi yang tersedia: count, nunique, sum, mean,
                            min, max. count dapat memakai column "*".
                            Karena survey_data berisi jawaban vertikal dan ID dapat
                            berulang untuk multiple_answer, gunakan fungsi nunique
                            pada kolom ID bila yang diminta adalah jumlah responden
                            unik. Gunakan count "*" bila yang diminta jumlah baris
                            atau jumlah jawaban.
                            Isi sort_by hanya memakai kolom hasil yang tersedia.
                            limit wajib 1 sampai 5000. Untuk permintaan tabel tanpa
                            agregasi, isi columns dan kosongkan group_by serta
                            aggregations. Gunakan action "table" untuk menampilkan
                            tabel saja. Hanya jika pengguna secara jelas meminta Excel,
                            file, ekspor, atau unduhan, gunakan action "excel"; aplikasi
                            akan menampilkan tabel dan tombol unduh .xlsx. Jika pengguna
                            meminta data dari kedua sumber, sertakan query database dan
                            query conversation dalam daftar queries agar Excel memiliki
                            satu sheet untuk tiap hasil. Jangan sertakan DATA_ACTION
                            jika pengguna tidak meminta tabel, tabulasi, ekstraksi, atau
                            file Excel. Aksi DATA_ACTION dapat disertakan bersama
                            VISUAL_ACTION, dan kedua blok JSON harus valid.
                            """
                            client = genai.Client(api_key=st.secrets["GEMINI_API_KEY"])
                            with st.spinner("Sedang melakukan analisis..."):
                                response = client.models.generate_content(
                                    model="gemini-3.8-flash",
                                    contents=ai_prompt,
                                )
                            if not response.text:
                                chat_error = "AI tidak mengembalikan jawaban."
                            else:
                                response_without_data_action = response.text
                                data_action = None
                                try:
                                    (
                                        response_without_data_action,
                                        data_action,
                                    ) = parse_data_action(response.text)
                                except ValueError as error:
                                    response_without_data_action = remove_action_block(
                                        response.text,
                                        "||DATA_ACTION||",
                                    )
                                    chat_error = (
                                        f"Tidak dapat mengambil data: {error}"
                                    )
                                try:
                                    answer_text, visual_action = parse_visual_action(
                                        response_without_data_action
                                    )
                                except ValueError as error:
                                    answer_text = response_without_data_action.partition(
                                        "||VISUAL_ACTION||"
                                    )[0].strip()
                                    visual_action = None
                                    chat_error = chat_error or (
                                        f"Tidak dapat membuat visualisasi: {error}"
                                    )
                                assistant_message: dict[str, object] = {
                                    "role": "assistant",
                                    "content": answer_text,
                                }
                                if visual_action is not None:
                                    assistant_message["visual_action"] = visual_action
                                if data_action is not None:
                                    assistant_message["data_action"] = data_action
                                messages.append(assistant_message)
                                chat_histories[selected_database] = messages
                                try:
                                    save_chat_histories(
                                        st.session_state.authenticated_username,
                                        chat_histories,
                                    )
                                except OSError as error:
                                    chat_error = (
                                        "Jawaban tampil tetapi gagal disimpan ke "
                                        f"riwayat chat: {error}"
                                    )

                with chat_container:
                    for message_index, message in enumerate(messages):
                        message_role = message.get("role")
                        message_content = message.get("content")
                        if (
                            message_role not in {"user", "assistant"}
                            or not isinstance(message_content, str)
                        ):
                            continue
                        with st.chat_message(message_role):
                            st.markdown(message_content)
                            visual_action = message.get("visual_action")
                            if (
                                message["role"] == "assistant"
                                and isinstance(visual_action, dict)
                            ):
                                render_visual_action(
                                    visual_action,
                                    tren_jawaban,
                                    message_index,
                                )
                            data_action = message.get("data_action")
                            if (
                                message["role"] == "assistant"
                                and isinstance(data_action, dict)
                            ):
                                render_data_action(
                                    data_action,
                                    database_tables,
                                    build_conversation_dataframe(messages),
                                    message_index,
                                )
                    if chat_error:
                        st.error(chat_error)

# ==============================================================================
# TAB 5: ANALISIS STATISTIK
# ==============================================================================
if selected_page == "statistics":
    username = st.session_state.authenticated_username
    try:
        statistical_conversations = load_statistics_histories(username)
    except (OSError, json.JSONDecodeError, ValueError) as error:
        st.error(f"Gagal membaca riwayat analisis statistik: {error}")
    else:
        selected_conversation_id = st.session_state.get(
            "selected_statistics_conversation"
        )
        conversation_ids = [
            str(conversation["id"]) for conversation in statistical_conversations
        ]
        if selected_conversation_id not in conversation_ids:
            selected_conversation_id = None

        if st.button(
            "Percakapan baru",
            icon=":material/add:",
            type="primary",
            key="new_statistics_conversation",
        ):
            now = datetime.now().astimezone().isoformat(timespec="seconds")
            new_conversation = {
                "id": uuid.uuid4().hex,
                "title": "Analisis baru",
                "created_at": now,
                "updated_at": now,
                "dataset_filename": None,
                "messages": [],
            }
            statistical_conversations.insert(0, new_conversation)
            try:
                save_statistics_histories(username, statistical_conversations)
            except OSError as error:
                st.error(f"Gagal membuat percakapan baru: {error}")
            else:
                st.session_state.selected_statistics_conversation = (
                    new_conversation["id"]
                )
                st.rerun()

        if not statistical_conversations:
            now = datetime.now().astimezone().isoformat(timespec="seconds")
            initial_conversation = {
                "id": uuid.uuid4().hex,
                "title": "Analisis baru",
                "created_at": now,
                "updated_at": now,
                "dataset_filename": None,
                "messages": [],
            }
            statistical_conversations.append(initial_conversation)
            try:
                save_statistics_histories(username, statistical_conversations)
            except OSError as error:
                st.error(f"Gagal menyiapkan riwayat percakapan: {error}")
                st.stop()
            st.session_state.selected_statistics_conversation = (
                initial_conversation["id"]
            )
            st.rerun()

        conversation_titles = {
            str(conversation["id"]): (
                f"{conversation['title']} · "
                f"{conversation['updated_at'][:16].replace('T', ' ')}"
            )
            for conversation in statistical_conversations
        }
        if st.session_state.get("selected_statistics_conversation") not in conversation_titles:
            st.session_state.selected_statistics_conversation = str(
                statistical_conversations[0]["id"]
            )

        selected_conversation_id = st.selectbox(
            "Buka riwayat percakapan",
            options=list(conversation_titles),
            format_func=lambda conversation_id: conversation_titles[conversation_id],
            key="selected_statistics_conversation",
        )
        selected_conversation = next(
            conversation
            for conversation in statistical_conversations
            if conversation["id"] == selected_conversation_id
        )

        confirm_delete_history = st.checkbox(
            "Saya yakin ingin menghapus percakapan ini dan file CSV-nya.",
            key=f"confirm_delete_statistics_{selected_conversation_id}",
        )
        if st.button(
            "Hapus percakapan",
            icon=":material/delete:",
            disabled=not confirm_delete_history,
            key=f"delete_statistics_{selected_conversation_id}",
        ):
            try:
                statistical_conversations = delete_statistics_conversation(
                    username,
                    statistical_conversations,
                    selected_conversation_id,
                )
            except (FileNotFoundError, OSError, ValueError) as error:
                st.error(f"Gagal menghapus percakapan: {error}")
            else:
                st.session_state.pop("selected_statistics_conversation", None)
                st.rerun()

        st.divider()
        uploader_key = f"statistics_csv_{selected_conversation_id}"
        uploaded_csv = st.file_uploader(
            "Unggah data CSV",
            type=["csv"],
            key=uploader_key,
            help="CSV disimpan di folder riwayat pribadi agar dapat dibuka kembali.",
        )
        if st.button(
            "Gunakan CSV untuk analisis",
            disabled=uploaded_csv is None,
            key=f"load_statistics_csv_{selected_conversation_id}",
        ):
            try:
                csv_bytes = uploaded_csv.getvalue()
                uploaded_dataframe = pd.read_csv(io.BytesIO(csv_bytes))
                if uploaded_dataframe.empty or len(uploaded_dataframe.columns) == 0:
                    raise ValueError("File CSV tidak berisi data atau nama kolom.")
                dataset_path = get_statistics_dataset_path(
                    username,
                    selected_conversation_id,
                )
                dataset_path.parent.mkdir(parents=True, exist_ok=True)
                dataset_path.write_bytes(csv_bytes)
                selected_conversation["dataset_filename"] = (
                    f"{selected_conversation_id}.csv"
                )
                selected_conversation["title"] = uploaded_csv.name
                selected_conversation["updated_at"] = (
                    datetime.now()
                    .astimezone()
                    .isoformat(timespec="seconds")
                )
                save_statistics_histories(username, statistical_conversations)
            except (
                OSError,
                UnicodeDecodeError,
                pd.errors.ParserError,
                pd.errors.EmptyDataError,
                ValueError,
            ) as error:
                st.error(f"CSV tidak dapat digunakan: {error}")
            else:
                st.rerun()

        statistics_dataframe = None
        dataset_filename = selected_conversation.get("dataset_filename")
        if isinstance(dataset_filename, str):
            dataset_path = get_statistics_dataset_path(
                username,
                selected_conversation_id,
            )
            try:
                statistics_dataframe = pd.read_csv(dataset_path)
            except FileNotFoundError:
                st.warning(
                    "File CSV untuk percakapan ini tidak ditemukan. Unggah kembali "
                    "CSV untuk melanjutkan analisis."
                )
            except (
                OSError,
                UnicodeDecodeError,
                pd.errors.ParserError,
                pd.errors.EmptyDataError,
            ) as error:
                st.error(f"Gagal membuka file CSV riwayat: {error}")

        analysis_result = None
        analysis_context = "Belum ada file CSV yang diunggah."
        if statistics_dataframe is not None:
            st.write(
                f"**Dataset aktif:** `{selected_conversation['title']}` · "
                f"{len(statistics_dataframe):,} baris × "
                f"{len(statistics_dataframe.columns):,} kolom"
            )
            st.dataframe(statistics_dataframe.head(20), hide_index=True)
            try:
                analysis_result = analyze_dataframe_automatically(
                    statistics_dataframe
                )
            except (ValueError, TypeError, FloatingPointError) as error:
                st.error(f"Analisis statistik otomatis gagal: {error}")
            else:
                st.subheader("Profil kolom")
                st.dataframe(
                    analysis_result["column_profile"],
                    hide_index=True,
                )

                numeric_summary = analysis_result["numeric_summary"]
                if not numeric_summary.empty:
                    st.subheader("Statistik deskriptif numerik")
                    st.dataframe(numeric_summary, hide_index=True)

                categorical_frequencies = analysis_result[
                    "categorical_frequency"
                ]
                if categorical_frequencies:
                    st.subheader("Distribusi kategori")
                    for column_name, frequency_dataframe in (
                        categorical_frequencies.items()
                    ):
                        st.markdown(f"**{column_name}**")
                        st.dataframe(frequency_dataframe, hide_index=True)

                st.subheader("Uji statistik otomatis")
                test_results = analysis_result["tests"]
                if test_results.empty:
                    st.info(
                        "Belum cukup data untuk uji inferensial. Pastikan terdapat "
                        "variabel numerik dengan variasi dan data yang memadai."
                    )
                else:
                    st.dataframe(test_results, hide_index=True)
                    st.caption(
                        "Uji bersifat eksploratif; p-value belum dikoreksi untuk "
                        "pengujian berganda dan tidak membuktikan hubungan sebab-akibat."
                    )

                numeric_context = (
                    numeric_summary.head(20).to_string(index=False)
                    if not numeric_summary.empty
                    else "Tidak ada kolom numerik."
                )
                tests_context = (
                    test_results.head(30).to_string(index=False)
                    if not test_results.empty
                    else "Belum ada hasil uji inferensial."
                )
                analysis_context = (
                    f"Ukuran data: {statistics_dataframe.shape[0]} baris, "
                    f"{statistics_dataframe.shape[1]} kolom.\n"
                    f"Tipe kolom:\n{statistics_dataframe.dtypes.to_string()}\n"
                    f"Sampel lima baris:\n"
                    f"{statistics_dataframe.head(5).to_string(max_cols=20, max_colwidth=60)}\n"
                    f"Ringkasan numerik:\n{numeric_context}\n"
                    f"Hasil uji otomatis:\n{tests_context}"
                )

        st.divider()
        st.subheader("Ruang Konsultasi & Olah Data Statistik")
        st.caption(
            "Pertanyaan dapat membahas pengolahan data, metode statistik, maupun "
            "referensi umum. Jika CSV aktif, ringkasan dan sampel data akan dikirim "
            "ke Gemini sebagai konteks. Untuk permintaan uji, Gemini memilih metode "
            "dan variabel; perhitungan dijalankan lokal menggunakan uji yang diizinkan."
        )
        messages = selected_conversation["messages"]
        prompt = st.text_input(
            "Tulis pertanyaan atau analisis statistik yang diinginkan:",
            placeholder=(
                "Contoh: Apakah nilai score berbeda antara kelompok A dan B? "
                "Tolong uji signifikansinya."
            ),
            max_chars=2000,
            key=f"statistics_prompt_{selected_conversation_id}",
        )
        if st.button(
            "Jalankan Analisis AI",
            type="primary",
            key=f"run_statistics_ai_{selected_conversation_id}",
        ):
            if not prompt.strip():
                st.warning("Silakan ketik pertanyaan atau instruksi analisis terlebih dahulu.")
            else:
                prompt = prompt.strip()
                messages.append({"role": "user", "content": prompt})
                if selected_conversation["title"] == "Analisis baru":
                    selected_conversation["title"] = prompt[:60]
                selected_conversation["updated_at"] = (
                    datetime.now().astimezone().isoformat(timespec="seconds")
                )
                try:
                    save_statistics_histories(username, statistical_conversations)
                except OSError as error:
                    st.error(f"Pesan terkirim tetapi gagal menyimpan riwayat: {error}")
                else:
                    system_instruction = (
                        "Anda adalah asisten ahli statistik dan pengolahan data. "
                        "Jawab dalam Bahasa Indonesia. Anda dapat membantu pertanyaan "
                        "umum dan metode. Jika diminta referensi, jangan mengarang "
                        "judul, DOI, URL, atau kutipan, dan sarankan verifikasi detail "
                        "bibliografi. Untuk setiap permintaan analisis data, jangan "
                        "menghitung hasil atau menghasilkan kode. Pilih metode dan "
                        "kolom, lalu keluarkan hanya JSON valid dengan salah satu "
                        "bentuk berikut: "
                        '{"intent":"general","answer":"jawaban umum"} atau '
                        '{"intent":"analysis","test":"chi_square|pearson|spearman|'
                        'independent_t|mann_whitney|one_way_anova|kruskal",'
                        '"columns":["nama kolom"],"group_column":null,'
                        '"rationale":"alasan singkat pemilihan uji"}. '
                        "Chi-Square memerlukan dua kolom kategori; Pearson/Spearman "
                        "dua kolom numerik; independent_t dan mann_whitney satu kolom "
                        "numerik dalam kolom groups dua kategori; one_way_anova dan "
                        "kruskal satu kolom numerik dengan group_column kategori. "
                        "Untuk semua analisis, isi columns dengan nama kolom persis "
                        "seperti skema. Hanya pilih metode jika cocok dengan permintaan "
                        "dan tipe data. Jangan sertakan kode Python."
                    )
                    history_context = "\n".join(
                        f"{'Pengguna' if message['role'] == 'user' else 'Asisten'}: "
                        f"{message['content']}"
                        for message in messages[-21:-1]
                    )
                    ai_prompt = (
                        f"Konteks dataset dan hasil statistik:\n{analysis_context}\n\n"
                        f"Riwayat percakapan:\n"
                        f"{history_context or '(Belum ada percakapan sebelumnya)'}\n\n"
                        f"Instruksi terbaru:\n{prompt}"
                    )
                    try:
                        client = genai.Client(api_key=st.secrets["GEMINI_API_KEY"])
                        with st.spinner("Gemini sedang meneliti variabel dan memilih metode..."):
                            response = client.models.generate_content(
                                model="gemini-3.8-flash",
                                contents=ai_prompt,
                                config=types.GenerateContentConfig(
                                    system_instruction=system_instruction,
                                    temperature=0.1,
                                ),
                            )
                        if not response.text:
                            st.error("Gemini tidak mengembalikan jawaban.")
                        else:
                            raw_plan = response.text.strip()
                            if raw_plan.startswith("```"):
                                raw_plan = raw_plan.partition("\n")[2]
                                raw_plan = raw_plan.rsplit("```", 1)[0].strip()
                            analysis_plan = json.loads(raw_plan)
                            if not isinstance(analysis_plan, dict):
                                raise ValueError("Format rencana analisis dari Gemini tidak valid.")

                            if analysis_plan.get("intent") == "general":
                                answer_text = analysis_plan.get("answer")
                                if not isinstance(answer_text, str) or not answer_text.strip():
                                    raise ValueError("Gemini tidak mengembalikan jawaban umum.")
                            elif analysis_plan.get("intent") == "analysis":
                                if statistics_dataframe is None:
                                    raise ValueError(
                                        "Unggah CSV terlebih dahulu untuk menjalankan uji statistik."
                                    )
                                test_result = run_requested_statistical_test(
                                    statistics_dataframe,
                                    analysis_plan,
                                )
                                rationale = analysis_plan.get("rationale")
                                result_lines = [
                                    f"### {test_result['test']}",
                                    f"- **Variabel:** {', '.join(test_result['columns'])}",
                                    f"- **Jumlah data (N):** {test_result['n']}",
                                    f"- **Statistik uji:** {test_result['statistic']:.6g}",
                                    f"- **p-value:** {test_result['p_value']:.6g}",
                                ]
                                if test_result["group_column"]:
                                    result_lines.append(
                                        f"- **Kelompok:** {test_result['group_column']}"
                                    )
                                if test_result["extra"]:
                                    result_lines.append(f"- {test_result['extra']}")
                                result_lines.append(
                                    "- **Interpretasi awal:** "
                                    + (
                                        "p < 0,05; hasil menunjukkan indikasi statistik "
                                        "yang signifikan pada ambang 5%."
                                        if test_result["p_value"] < 0.05
                                        else "p >= 0,05; belum ada bukti statistik yang "
                                        "cukup pada ambang 5%."
                                    )
                                )
                                if isinstance(rationale, str) and rationale.strip():
                                    result_lines.append(
                                        f"\n**Alasan pemilihan uji:** {rationale.strip()}"
                                    )
                                result_lines.append(
                                    "\nHasil ini bersifat statistik/asosiatif dan "
                                    "tidak dengan sendirinya membuktikan sebab-akibat. "
                                    "Periksa asumsi uji dan konteks penelitian sebelum "
                                    "menarik kesimpulan."
                                )
                                answer_text = "\n".join(result_lines)
                            else:
                                raise ValueError(
                                    "Gemini tidak mengembalikan intent general atau analysis."
                                )

                            messages.append(
                                {"role": "assistant", "content": answer_text}
                            )
                            selected_conversation["updated_at"] = (
                                datetime.now().astimezone().isoformat(timespec="seconds")
                            )
                            save_statistics_histories(
                                username,
                                statistical_conversations,
                            )
                    except (
                        KeyError,
                        OSError,
                        ValueError,
                        json.JSONDecodeError,
                        genai.errors.APIError,
                    ) as error:
                        st.error(f"Gagal menjalankan konsultasi analisis: {error}")

        for message in messages:
            with st.chat_message(str(message["role"])):
                st.markdown(str(message["content"]))

# ==============================================================================
# TAB 4: EDIT FOTO
# ==============================================================================
if selected_page == "photo":
    with st.form("photo_generation_form"):
        photo_prompt = st.text_area(
            "Prompt",
            placeholder="Contoh: Buatkan cover PPT dengan judul x dan tema y.",
            max_chars=4000,
        )
        uploaded_photo = st.file_uploader(
            "Foto sumber (opsional)",
            type=["png", "jpg", "jpeg", "webp"],
            help="Kosongkan jika ingin membuat gambar hanya dari prompt.",
            max_upload_size=10,
        )
        photo_submitted = st.form_submit_button(
            "Buat / edit foto",
            type="primary",
        )

    if photo_submitted:
        cleaned_prompt = photo_prompt.strip()
        if not cleaned_prompt:
            st.error("Prompt wajib diisi.")
        elif "GEMINI_API_KEY" not in st.secrets:
            st.error("API Key belum diset di Secrets Streamlit.")
        else:
            request_contents: list[str | types.Part] = [
                types.Part.from_text(text=cleaned_prompt)
            ]
            upload_error = None
            if uploaded_photo is not None:
                photo_bytes = uploaded_photo.getvalue()
                if len(photo_bytes) > 10 * 1024 * 1024:
                    upload_error = "Ukuran foto sumber maksimal 10 MB."
                else:
                    try:
                        with Image.open(BytesIO(photo_bytes)) as source_image:
                            source_format = source_image.format
                            source_image.verify()
                        mime_type_by_format = {
                            "PNG": "image/png",
                            "JPEG": "image/jpeg",
                            "WEBP": "image/webp",
                        }
                        source_mime_type = mime_type_by_format.get(source_format)
                        if source_mime_type is None:
                            upload_error = "Format foto sumber harus PNG, JPEG, atau WebP."
                        else:
                            request_contents.append(
                                types.Part.from_bytes(
                                    data=photo_bytes,
                                    mime_type=source_mime_type,
                                )
                            )
                    except (
                        Image.DecompressionBombError,
                        UnidentifiedImageError,
                        OSError,
                        ValueError,
                    ):
                        upload_error = "File yang diunggah bukan gambar yang valid."

            if upload_error:
                st.error(upload_error)
            else:
                try:
                    client = genai.Client(api_key=st.secrets["GEMINI_API_KEY"])
                    with st.spinner("Sedang membuat atau mengedit foto..."):
                        image_response = client.models.generate_content(
                            model="gemini-3.8-flash-image",
                            contents=request_contents,
                            config=types.GenerateContentConfig(
                                response_modalities=["TEXT", "IMAGE"]
                            ),
                        )
                except genai.errors.APIError as error:
                    st.error(f"Layanan Gemini gagal menghasilkan foto: {error}")
                else:
                    generated_images: list[tuple[bytes, str]] = []
                    response_text = ""
                    for candidate in image_response.candidates or []:
                        if candidate.content is None:
                            continue
                        for part in candidate.content.parts or []:
                            if part.text:
                                response_text = part.text
                            if (
                                part.inline_data is not None
                                and isinstance(part.inline_data.data, bytes)
                            ):
                                generated_images.append(
                                    (
                                        part.inline_data.data,
                                        part.inline_data.mime_type or "",
                                    )
                                )

                    if not generated_images:
                        st.error(
                            response_text
                            or "Gemini tidak mengembalikan gambar. Coba ubah prompt."
                        )
                    else:
                        saved_count = 0
                        for image_bytes, mime_type in generated_images:
                            try:
                                with Image.open(BytesIO(image_bytes)) as output_image:
                                    output_image.verify()
                                save_generated_photo(
                                    st.session_state.authenticated_username,
                                    cleaned_prompt,
                                    image_bytes,
                                    mime_type,
                                )
                            except (
                                Image.DecompressionBombError,
                                UnidentifiedImageError,
                                OSError,
                                ValueError,
                            ) as error:
                                st.error(f"Gagal menyimpan hasil foto: {error}")
                            else:
                                saved_count += 1
                        if saved_count:
                            st.success(f"{saved_count} foto berhasil disimpan.")
                        if response_text:
                            st.caption(response_text)

    st.subheader("Riwayat foto")
    try:
        photo_history = load_photo_history(
            st.session_state.authenticated_username
        )
    except (OSError, json.JSONDecodeError, ValueError) as error:
        st.error(f"Gagal membaca riwayat foto: {error}")
    else:
        if not photo_history:
            st.info("Belum ada foto yang tersimpan.")
        for entry in photo_history:
            st.markdown(f"**{entry['created_at']}**")
            st.caption(entry["prompt"])
            photo_path = (
                get_photo_history_directory(st.session_state.authenticated_username)
                / entry["filename"]
            )
            try:
                saved_photo_bytes = photo_path.read_bytes()
            except OSError as error:
                st.error(f"Foto riwayat tidak dapat dibaca: {error}")
                continue

            st.image(
                saved_photo_bytes,
                alt="Foto dari riwayat pengguna",
                width="stretch",
            )
            with st.container(horizontal=True):
                st.download_button(
                    "Unduh",
                    data=saved_photo_bytes,
                    file_name=f"statlabchat_{entry['filename']}",
                    mime=entry["mime_type"],
                    key=f"download_photo_{entry['id']}",
                )
                if st.button(
                    "Hapus",
                    key=f"delete_photo_{entry['id']}",
                    icon=":material/delete:",
                ):
                    try:
                        delete_generated_photo(
                            st.session_state.authenticated_username,
                            entry["id"],
                        )
                    except (FileNotFoundError, OSError, ValueError) as error:
                        st.error(f"Gagal menghapus foto: {error}")
                    else:
                        st.rerun()
